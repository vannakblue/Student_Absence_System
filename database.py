"""
ប្រព័ន្ធគ្រប់គ្រងអវត្តមានសិស្ស និងគ្រូបង្រៀន (Student & Teacher Absence Management System)
Database Layer with SQLite
"""

import sqlite3
import os
import sys
from datetime import datetime, timedelta
from werkzeug.security import generate_password_hash, check_password_hash

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "attendance.db")


def get_db_connection():
    """បង្កើត Connection ទៅកាន់ SQLite Database"""
    conn = sqlite3.connect(DB_PATH, timeout=30.0)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON;")
    conn.execute("PRAGMA journal_mode = WAL;")
    return conn


def init_db():
    """បង្កើត Tables ទាំងអស់ប្រសិនបើមិនទាន់មាន"""
    conn = get_db_connection()
    cursor = conn.cursor()

    # 1. Classes Table (ថ្នាក់រៀន)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS classes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            class_name TEXT NOT NULL UNIQUE,
            grade_level INTEGER NOT NULL,
            shift TEXT NOT NULL DEFAULT 'Morning', -- Morning (ព្រឹក), Afternoon (រសៀល), FullDay
            room_number TEXT,
            academic_year TEXT DEFAULT '2026-2027',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
    """)

    # 2. Teachers Table (គ្រូបង្រៀន)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS teachers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_code TEXT NOT NULL UNIQUE,
            full_name_kh TEXT NOT NULL,
            full_name_en TEXT,
            gender TEXT NOT NULL, -- ប្រុស (M) / ស្រី (F)
            phone TEXT,
            email TEXT,
            subject TEXT NOT NULL, -- មុខវិជ្ជាឯកទេស
            total_hours_weekly INTEGER DEFAULT 18, -- បន្ទុកម៉ោងបង្រៀនក្នុងមួយសប្តាហ៍
            status TEXT DEFAULT 'Active', -- Active, Inactive, OnLeave
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
    """)

    # 3. Students Table (សិស្សានុសិស្ស)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            student_code TEXT NOT NULL UNIQUE,
            full_name_kh TEXT NOT NULL,
            full_name_en TEXT,
            gender TEXT NOT NULL, -- ប្រុស (M) / ស្រី (F)
            dob TEXT, -- ថ្ងៃខែឆ្នាំកំណើត YYYY-MM-DD
            class_id INTEGER NOT NULL,
            parent_name TEXT,
            parent_phone TEXT,
            status TEXT DEFAULT 'Active', -- Active, Inactive, Transferred
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (class_id) REFERENCES classes (id) ON DELETE RESTRICT
        );
    """)

    # 4. Teacher Attendance Table (ស្រង់វត្តមានគ្រូបង្រៀន)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS teacher_attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT NOT NULL, -- YYYY-MM-DD
            teacher_id INTEGER NOT NULL,
            shift TEXT NOT NULL, -- Morning, Afternoon
            period TEXT NOT NULL, -- ម៉ោងទី 1-6 ឬ 1-12
            status TEXT NOT NULL, -- PRESENT (វត្តមាន), PERMISSION (ច្បាប់), ABSENT (អវត្តមាន), LATE (យឺត)
            reason TEXT,
            substitute_teacher_id INTEGER,
            notes TEXT,
            recorded_by TEXT DEFAULT 'Admin',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES teachers (id) ON DELETE CASCADE,
            FOREIGN KEY (substitute_teacher_id) REFERENCES teachers (id) ON DELETE SET NULL,
            UNIQUE (date, teacher_id, shift, period)
        );
    """)

    # 5. Student Attendance Table (ស្រង់វត្តមានសិស្សានុសិស្ស)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS student_attendance (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT NOT NULL, -- YYYY-MM-DD
            student_id INTEGER NOT NULL,
            class_id INTEGER NOT NULL,
            shift TEXT NOT NULL, -- Morning, Afternoon, FullDay
            period TEXT DEFAULT 'Daily', -- Daily ឬ ម៉ោងទី
            status TEXT NOT NULL, -- PRESENT, PERMISSION, ABSENT, LATE
            reason TEXT,
            recorded_by TEXT DEFAULT 'Teacher',
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (student_id) REFERENCES students (id) ON DELETE CASCADE,
            FOREIGN KEY (class_id) REFERENCES classes (id) ON DELETE CASCADE,
            UNIQUE (date, student_id, shift, period)
        );
    """)

    # 6. Leave Requests Table (ពាក្យសុំច្បាប់)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS leave_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            person_type TEXT NOT NULL, -- TEACHER, STUDENT
            person_id INTEGER NOT NULL,
            start_date TEXT NOT NULL,
            end_date TEXT NOT NULL,
            reason TEXT NOT NULL,
            status TEXT DEFAULT 'Approved', -- Pending, Approved, Rejected
            approved_by TEXT DEFAULT 'Headmaster',
            notes TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
    """)

    # 7. System Settings & Google Sheets Configuration (ការកំណត់ប្រព័ន្ធ)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
    """)

    # 8. Timetable Slots Table (កាលវិភាគបង្រៀន)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS timetable_slots (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            class_id INTEGER,
            class_code TEXT NOT NULL,
            day_code TEXT NOT NULL, -- ច, អ, ព, ព្រ, សុ, ស
            day_name TEXT NOT NULL, -- ចន្ទ, អង្គារ, ពុធ, ព្រហស្បតិ៍, សុក្រ, សៅរ៍
            period_num INTEGER NOT NULL, -- 1 ដល់ 8
            shift TEXT NOT NULL, -- Morning, Afternoon
            subject_code TEXT,
            subject_name TEXT,
            teacher_id INTEGER,
            teacher_code TEXT,
            teacher_name TEXT,
            room_number TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (class_id) REFERENCES classes (id) ON DELETE CASCADE,
            FOREIGN KEY (teacher_id) REFERENCES teachers (id) ON DELETE SET NULL,
            UNIQUE (class_code, day_code, period_num)
        );
    """)

    # Ensure room_number column exists in timetable_slots
    try:
        cursor.execute("ALTER TABLE timetable_slots ADD COLUMN room_number TEXT;")
    except Exception:
        pass

    # Ensure classes table has homeroom_teacher_id and telegram_chat_id
    try:
        cursor.execute("ALTER TABLE classes ADD COLUMN homeroom_teacher_id INTEGER;")
    except Exception:
        pass
    try:
        cursor.execute("ALTER TABLE classes ADD COLUMN telegram_chat_id TEXT;")
    except Exception:
        pass

    # Ensure students table has parent_telegram
    try:
        cursor.execute("ALTER TABLE students ADD COLUMN parent_telegram TEXT;")
    except Exception:
        pass

    # 9. Subjects Table (មុខវិជ្ជាកម្មវិធីសិក្សា)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS subjects (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT NOT NULL UNIQUE,
            name_kh TEXT NOT NULL,
            weekly_hours_g7 INTEGER DEFAULT 0,
            weekly_hours_g8 INTEGER DEFAULT 0,
            weekly_hours_g9 INTEGER DEFAULT 0,
            weekly_hours_g10 INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );
    """)

    # 10. Users Table (គណនីប្រើប្រាស់ប្រព័ន្ធ)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            plain_password_hint TEXT,
            role TEXT NOT NULL DEFAULT 'teacher', -- 'admin', 'teacher'
            full_name_kh TEXT NOT NULL,
            teacher_id INTEGER,
            phone TEXT,
            is_active INTEGER DEFAULT 1,
            last_login TIMESTAMP,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES teachers (id) ON DELETE SET NULL
        );
    """)

    # 11. Attendance Audit Logs Table (កត់ត្រាការដាក់ស្នើវត្តមានសិស្ស និងតាមដានចន្លោះម៉ោង)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS attendance_audit_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT NOT NULL,
            class_id INTEGER NOT NULL,
            shift TEXT NOT NULL,
            period TEXT NOT NULL,
            period_num INTEGER,
            teacher_id INTEGER,
            submitted_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            phase TEXT NOT NULL, -- 'first_30', 'second_30'
            submission_count INTEGER DEFAULT 1,
            ip_address TEXT,
            notes TEXT
        );
    """)

    # Insert default settings if not exists
    default_settings = [
        ("school_name_kh", "វិទ្យាល័យ ហ៊ុន សែន (សាលាគំរូ)"),
        ("school_name_en", "Hun Sen Model High School"),
        ("academic_year", "2026-2027"),
        ("google_sheet_id", ""),
        ("google_credentials_file", "credentials.json"),
        ("auto_sync_sheets", "0"),
        ("last_sync_time", "Never"),
        ("telegram_bot_token", ""),
        ("telegram_admin_chat_id", ""),
        ("telegram_notify_absence", "1"),
        ("telegram_notify_homeroom", "1"),
        ("telegram_notify_parent", "1"),
    ]
    for key, val in default_settings:
        cursor.execute(
            "INSERT OR IGNORE INTO settings (key, value) VALUES (?, ?);", (key, val)
        )

    conn.commit()
    conn.close()
    print("Database initialized successfully at:", DB_PATH)


# ==========================================
# CLASSES CRUD
# ==========================================
def get_classes():
    conn = get_db_connection()
    classes = conn.execute("""
        SELECT c.*, 
               t.full_name_kh as homeroom_teacher_name,
               (SELECT COUNT(*) FROM students s WHERE s.class_id = c.id AND s.status = 'Active') as student_count
        FROM classes c
        LEFT JOIN teachers t ON c.homeroom_teacher_id = t.id
        ORDER BY c.grade_level ASC, c.class_name ASC
    """).fetchall()
    conn.close()
    return [dict(c) for c in classes]


def get_class_by_id(class_id):
    conn = get_db_connection()
    c = conn.execute("""
        SELECT c.*, t.full_name_kh as homeroom_teacher_name
        FROM classes c
        LEFT JOIN teachers t ON c.homeroom_teacher_id = t.id
        WHERE c.id = ?
    """, (class_id,)).fetchone()
    conn.close()
    return dict(c) if c else None


def update_class_homeroom_and_telegram(class_id, homeroom_teacher_id=None, telegram_chat_id=None):
    """កែប្រែគ្រូបន្ទុកថ្នាក់ និង Telegram Group របស់ថ្នាក់"""
    conn = get_db_connection()
    conn.execute("""
        UPDATE classes
        SET homeroom_teacher_id = ?, telegram_chat_id = ?
        WHERE id = ?
    """, (homeroom_teacher_id, telegram_chat_id, class_id))
    conn.commit()
    conn.close()
    return True


def add_class(class_name, grade_level, shift="Morning", room_number="", academic_year="2026-2027"):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO classes (class_name, grade_level, shift, room_number, academic_year)
        VALUES (?, ?, ?, ?, ?)
    """, (class_name, grade_level, shift, room_number, academic_year))
    conn.commit()
    new_id = cursor.lastrowid
    conn.close()
    return new_id


# ==========================================
# TEACHERS CRUD
# ==========================================
def get_teachers(active_only=True):
    conn = get_db_connection()
    query = "SELECT * FROM teachers"
    if active_only:
        query += " WHERE status = 'Active'"
    query += " ORDER BY full_name_kh ASC"
    teachers = conn.execute(query).fetchall()
    conn.close()
    return [dict(t) for t in teachers]


def get_teacher_by_id(teacher_id):
    conn = get_db_connection()
    t = conn.execute("SELECT * FROM teachers WHERE id = ?", (teacher_id,)).fetchone()
    conn.close()
    return dict(t) if t else None


def add_teacher(code, full_name_kh, gender, phone="", email="", subject="", full_name_en="", total_hours=18):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO teachers (teacher_code, full_name_kh, full_name_en, gender, phone, email, subject, total_hours_weekly)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (code, full_name_kh, full_name_en, gender, phone, email, subject, total_hours))
    conn.commit()
    new_id = cursor.lastrowid
    conn.close()

    # Automatically create user account for this teacher
    try:
        create_user(
            username=code,
            password="123456",
            role="teacher",
            full_name_kh=full_name_kh,
            teacher_id=new_id,
            phone=phone,
            plain_hint="123456"
        )
    except Exception:
        pass

    return new_id


def update_teacher(teacher_id, code, full_name_kh, gender, phone="", email="", subject="", full_name_en="", status="Active", total_hours=18):
    conn = get_db_connection()
    conn.execute("""
        UPDATE teachers 
        SET teacher_code = ?, full_name_kh = ?, full_name_en = ?, gender = ?, phone = ?, email = ?, subject = ?, status = ?, total_hours_weekly = ?
        WHERE id = ?
    """, (code, full_name_kh, full_name_en, gender, phone, email, subject, status, total_hours, teacher_id))
    # Update matching user account
    conn.execute("""
        UPDATE users
        SET username = ?, full_name_kh = ?, phone = ?
        WHERE teacher_id = ?
    """, (code, full_name_kh, phone, teacher_id))
    conn.commit()
    conn.close()


def delete_teacher(teacher_id):
    conn = get_db_connection()
    conn.execute("DELETE FROM users WHERE teacher_id = ?", (teacher_id,))
    conn.execute("DELETE FROM timetable_slots WHERE teacher_id = ?", (teacher_id,))
    conn.execute("DELETE FROM teachers WHERE id = ?", (teacher_id,))
    conn.commit()
    conn.close()


# ==========================================
# STUDENTS CRUD
# ==========================================
def get_students(class_id=None, active_only=True):
    conn = get_db_connection()
    params = []
    query = """
        SELECT s.*, c.class_name, c.grade_level, c.shift
        FROM students s
        JOIN classes c ON s.class_id = c.id
        WHERE 1=1
    """
    if active_only:
        query += " AND s.status = 'Active'"
    if class_id:
        query += " AND s.class_id = ?"
        params.append(class_id)
    query += " ORDER BY c.class_name ASC, s.full_name_kh ASC"
    students = conn.execute(query, params).fetchall()
    conn.close()
    return [dict(s) for s in students]


def get_student_by_id(student_id):
    conn = get_db_connection()
    s = conn.execute("""
        SELECT s.*, c.class_name 
        FROM students s
        JOIN classes c ON s.class_id = c.id
        WHERE s.id = ?
    """, (student_id,)).fetchone()
    conn.close()
    return dict(s) if s else None


def add_student(code, full_name_kh, gender, class_id, dob="", full_name_en="", parent_name="", parent_phone="", parent_telegram=""):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO students (student_code, full_name_kh, full_name_en, gender, dob, class_id, parent_name, parent_phone, parent_telegram)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (code, full_name_kh, full_name_en, gender, dob, class_id, parent_name, parent_phone, parent_telegram))
    conn.commit()
    new_id = cursor.lastrowid
    conn.close()
    return new_id


def update_student(student_id, code, full_name_kh, gender, class_id, dob="", full_name_en="", parent_name="", parent_phone="", status="Active", parent_telegram=""):
    conn = get_db_connection()
    conn.execute("""
        UPDATE students
        SET student_code = ?, full_name_kh = ?, full_name_en = ?, gender = ?, dob = ?, class_id = ?, parent_name = ?, parent_phone = ?, status = ?, parent_telegram = ?
        WHERE id = ?
    """, (code, full_name_kh, full_name_en, gender, dob, class_id, parent_name, parent_phone, status, parent_telegram, student_id))
    conn.commit()
    conn.close()


def delete_student(student_id):
    conn = get_db_connection()
    conn.execute("DELETE FROM students WHERE id = ?", (student_id,))
    conn.commit()
    conn.close()


# ==========================================
# TEACHER ATTENDANCE OPERATIONS
# ==========================================
def get_teacher_attendance(date_str, shift="Morning", period="Session 1"):
    """ទាញយកវត្តមានគ្រូទាំងអស់តាមកាលបរិច្ឆេទ វេន និងម៉ោង"""
    conn = get_db_connection()
    rows = conn.execute("""
        SELECT t.id as teacher_id, t.teacher_code, t.full_name_kh, t.gender, t.subject, t.phone,
               ta.id as attendance_id,
               COALESCE(ta.status, 'PRESENT') as status,
               ta.reason,
               ta.substitute_teacher_id,
               sub.full_name_kh as substitute_name_kh,
               ta.notes
        FROM teachers t
        LEFT JOIN teacher_attendance ta 
               ON t.id = ta.teacher_id 
              AND ta.date = ? 
              AND ta.shift = ? 
              AND ta.period = ?
        LEFT JOIN teachers sub 
               ON ta.substitute_teacher_id = sub.id
        WHERE t.status = 'Active'
        ORDER BY t.full_name_kh ASC
    """, (date_str, shift, period)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def save_teacher_attendance(date_str, shift, period, records, recorded_by="Admin"):
    """
    រក្សាទុកកំណត់ត្រាវត្តមានគ្រូ
    records: list of dicts: [{'teacher_id': 1, 'status': 'PRESENT', 'reason': '', 'substitute_teacher_id': None, 'notes': ''}]
    """
    conn = get_db_connection()
    cursor = conn.cursor()
    for rec in records:
        teacher_id = rec.get("teacher_id")
        status = rec.get("status", "PRESENT")
        reason = rec.get("reason", "")
        sub_id = rec.get("substitute_teacher_id")
        if not sub_id or sub_id == "" or sub_id == "0":
            sub_id = None
        notes = rec.get("notes", "")

        cursor.execute("""
            INSERT INTO teacher_attendance (date, teacher_id, shift, period, status, reason, substitute_teacher_id, notes, recorded_by)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(date, teacher_id, shift, period) DO UPDATE SET
                status = excluded.status,
                reason = excluded.reason,
                substitute_teacher_id = excluded.substitute_teacher_id,
                notes = excluded.notes,
                recorded_by = excluded.recorded_by,
                created_at = CURRENT_TIMESTAMP
        """, (date_str, teacher_id, shift, period, status, reason, sub_id, notes, recorded_by))
    conn.commit()
    conn.close()
    return True


# ==========================================
# STUDENT ATTENDANCE OPERATIONS
# ==========================================
def get_student_attendance(class_id, date_str, shift="Morning", period="Daily"):
    """ទាញយកបញ្ជីវត្តមានសិស្សទាំងអស់ក្នុងថ្នាក់ តាមកាលបរិច្ឆេទ"""
    conn = get_db_connection()
    rows = conn.execute("""
        SELECT s.id as student_id, s.student_code, s.full_name_kh, s.full_name_en, s.gender, s.parent_phone, s.parent_name, s.parent_telegram,
               sa.id as attendance_id,
               COALESCE(sa.status, 'PRESENT') as status,
               sa.reason
        FROM students s
        LEFT JOIN student_attendance sa 
               ON s.id = sa.student_id 
              AND sa.date = ? 
              AND sa.shift = ? 
              AND sa.period = ?
        WHERE s.class_id = ? AND s.status = 'Active'
        ORDER BY s.full_name_kh ASC
    """, (date_str, shift, period, class_id)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def save_student_attendance(class_id, date_str, shift, period, records, recorded_by="Teacher"):
    """
    រក្សាទុកកំណត់ត្រាវត្តមានសិស្ស (Sparse Storage Model)៖
    - សិស្សអវត្តមាន (ABSENT, PERMISSION, LATE) ត្រូវរក្សាទុកក្នុង Database
    - សិស្សមានវត្តមាន (PRESENT) មិនត្រូវរក្សាទុកក្នុង Database ឡើយ (ហើយលុប record ចាស់បើធ្លាប់មាន)
    - កត់ត្រាក្នុង attendance_audit_logs ដើម្បីបញ្ជាក់ថាថ្នាក់នេះបានស្រង់រួចរាល់ ទោះបីគ្មានសិស្សអវត្តមានក៏ដោយ
    records: list of dicts: [{'student_id': 1, 'status': 'PRESENT', 'reason': ''}]
    """
    conn = get_db_connection()
    cursor = conn.cursor()
    for rec in records:
        student_id = rec.get("student_id")
        status = rec.get("status", "PRESENT")
        reason = rec.get("reason", "")

        if status in ("ABSENT", "PERMISSION", "LATE"):
            cursor.execute("""
                INSERT INTO student_attendance (date, student_id, class_id, shift, period, status, reason, recorded_by)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(date, student_id, shift, period) DO UPDATE SET
                    status = excluded.status,
                    reason = excluded.reason,
                    recorded_by = excluded.recorded_by,
                    created_at = CURRENT_TIMESTAMP
            """, (date_str, student_id, class_id, shift, period, status, reason, recorded_by))
        else:
            # PRESENT: remove any previous absence record for this slot
            cursor.execute("""
                DELETE FROM student_attendance
                WHERE date = ? AND student_id = ? AND shift = ? AND period = ?
            """, (date_str, student_id, shift, period))

    # Record in audit logs so we know this class/period was recorded even if 0 students were absent
    audit_check = cursor.execute("""
        SELECT COUNT(*) as cnt FROM attendance_audit_logs
        WHERE date = ? AND class_id = ? AND shift = ? AND period = ?
    """, (date_str, class_id, shift, str(period))).fetchone()
    if not audit_check or audit_check["cnt"] == 0:
        period_num = 1
        if "Session " in str(period):
            try:
                period_num = int(str(period).replace("Session ", "").strip())
            except Exception:
                pass
        cursor.execute("""
            INSERT INTO attendance_audit_logs (date, class_id, shift, period, period_num, teacher_id, phase, submission_count, notes)
            VALUES (?, ?, ?, ?, ?, NULL, 'first_30', 1, ?)
        """, (date_str, class_id, shift, str(period), period_num, f"Recorded by {recorded_by}"))

    conn.commit()
    conn.close()
    return True


# ==========================================
# LEAVE REQUESTS
# ==========================================
def get_leave_requests(person_type=None):
    conn = get_db_connection()
    query = """
        SELECT lr.*,
               CASE 
                   WHEN lr.person_type = 'TEACHER' THEN (SELECT full_name_kh FROM teachers WHERE id = lr.person_id)
                   WHEN lr.person_type = 'STUDENT' THEN (SELECT full_name_kh FROM students WHERE id = lr.person_id)
               END as person_name_kh,
               CASE 
                   WHEN lr.person_type = 'TEACHER' THEN (SELECT teacher_code FROM teachers WHERE id = lr.person_id)
                   WHEN lr.person_type = 'STUDENT' THEN (SELECT student_code FROM students WHERE id = lr.person_id)
               END as person_code
        FROM leave_requests lr
    """
    params = []
    if person_type:
        query += " WHERE lr.person_type = ?"
        params.append(person_type)
    query += " ORDER BY lr.start_date DESC, lr.id DESC"
    rows = conn.execute(query, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def add_leave_request(person_type, person_id, start_date, end_date, reason, approved_by="Headmaster", notes=""):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO leave_requests (person_type, person_id, start_date, end_date, reason, approved_by, notes)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (person_type, person_id, start_date, end_date, reason, approved_by, notes))
    conn.commit()
    new_id = cursor.lastrowid
    conn.close()
    return new_id


def get_teacher_teaching_days(teacher_id):
    """ទាញយកថ្ងៃនៃសប្តាហ៍ដែលគ្រូមានម៉ោងបង្រៀន (ច, អ, ព, ព្រ, សុ, ស)"""
    conn = get_db_connection()
    rows = conn.execute("""
        SELECT DISTINCT day_code, day_name
        FROM timetable_slots
        WHERE teacher_id = ?
        ORDER BY 
            CASE day_code
                WHEN 'ច' THEN 1
                WHEN 'អ' THEN 2
                WHEN 'ព' THEN 3
                WHEN 'ព្រ' THEN 4
                WHEN 'សុ' THEN 5
                WHEN 'ស' THEN 6
                ELSE 7
            END
    """, (teacher_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def validate_teacher_leave_eligibility(teacher_id, start_date, end_date=None, current_dt=None):
    """
    ផ្ទៀងផ្ទាត់លក្ខខណ្ឌសុំច្បាប់របស់គ្រូបង្រៀន៖
    1. មិនអនុញ្ញាតឱ្យសុំច្បាប់កាលបរិច្ឆេទអតីតកាល (start_date < today)
    2. ប្រសិនបើសុំសម្រាប់ថ្ងៃនេះ (start_date == today)៖ ត្រូវធ្វើឡើងមុនម៉ោងកំណត់ដោយ Admin (ឧ. 17:00 / ម៉ោង ៥ រសៀល)
    3. ថ្ងៃដែលសុំច្បាប់ ត្រូវតែជាថ្ងៃដែលគ្រូនោះមានម៉ោងបង្រៀនជាក់ស្តែងក្នុងកាលវិភាគ (Timetable)
    """
    if current_dt is None:
        current_dt = datetime.now()

    today_str = current_dt.strftime("%Y-%m-%d")
    current_time_str = current_dt.strftime("%H:%M")

    if not end_date:
        end_date = start_date

    try:
        start_dt = datetime.strptime(start_date, "%Y-%m-%d")
        end_dt = datetime.strptime(end_date, "%Y-%m-%d")
    except Exception:
        return False, "ទម្រង់កាលបរិច្ឆេទមិនត្រឹមត្រូវ (ត្រូវជា YYYY-MM-DD)", []

    if start_dt > end_dt:
        return False, "កាលបរិច្ឆេទចាប់ផ្តើមមិនអាចក្រោយកាលបរិច្ឆេទបញ្ចប់បានទេ", []

    # Check 1: មិនអនុញ្ញាតកាលបរិច្ឆេទអតីតកាល
    if start_date < today_str:
        return False, f"មិនអាចសុំច្បាប់សម្រាប់កាលបរិច្ឆេទកន្លងផុតទៅបានទេ (អាចសុំបានចាប់ពីថ្ងៃនេះ {today_str} ឡើងទៅ)!", []

    # Check 2: ប្រសិនបើសុំសម្រាប់ថ្ងៃនេះ ត្រូវតែមុនម៉ោងកំណត់ដោយ Admin (Cutoff Time)
    cutoff_time = get_setting("teacher_leave_cutoff_time", "17:00")
    if not cutoff_time or not str(cutoff_time).strip():
        cutoff_time = "17:00"
    cutoff_time = str(cutoff_time).strip()

    if start_date == today_str and current_time_str >= cutoff_time:
        return False, f"ផុតម៉ោងកំណត់សុំច្បាប់សម្រាប់ថ្ងៃនេះហើយ (កំណត់ត្រឹមម៉ោង {cutoff_time})! សូមទំនាក់ទំនងរដ្ឋបាលសាលាដោយផ្ទាល់។", []

    # Check 3: ពិនិត្យកាលវិភាគបង្រៀន (Timetable Slots)
    day_kh_map = {0: 'ច', 1: 'អ', 2: 'ព', 3: 'ព្រ', 4: 'សុ', 5: 'ស', 6: 'អា'}
    day_name_map = {0: 'ចន្ទ', 1: 'អង្គារ', 2: 'ពុធ', 3: 'ព្រហស្បតិ៍', 4: 'សុក្រ', 5: 'សៅរ៍', 6: 'អាទិត្យ'}

    teacher_slots = get_timetable_by_teacher(teacher_id)
    if not teacher_slots:
        return False, "លោកគ្រូ/អ្នកគ្រូ មិនទាន់មានកាលវិភាគបង្រៀននៅក្នុងប្រព័ន្ធនៅឡើយទេ! សូមទាក់ទងរដ្ឋបាល។", []

    cur = start_dt
    date_slots_map = {}
    non_teaching_dates = []

    while cur <= end_dt:
        cur_str = cur.strftime("%Y-%m-%d")
        w_idx = cur.weekday()
        if w_idx == 6:  # Sunday
            non_teaching_dates.append(f"{cur_str} (ថ្ងៃអាទិត្យ)")
            cur += timedelta(days=1)
            continue

        day_code = day_kh_map.get(w_idx, "")
        slots_on_day = [s for s in teacher_slots if s.get("day_code") == day_code]
        if slots_on_day:
            date_slots_map[cur_str] = slots_on_day
        else:
            non_teaching_dates.append(f"{cur_str} (ថ្ងៃ{day_name_map[w_idx]})")
        cur += timedelta(days=1)

    if not date_slots_map:
        days_str = ", ".join(non_teaching_dates)
        return False, f"លោកគ្រូ/អ្នកគ្រូ គ្មានម៉ោងបង្រៀននៅថ្ងៃ {days_str} ទេ! អាចសុំច្បាប់បានតែថ្ងៃដែលមានម៉ោងបង្រៀនប៉ុណ្ណោះ។", []

    all_affected_slots = []
    for d, s_list in date_slots_map.items():
        for s in s_list:
            all_affected_slots.append({
                "date": d,
                "class_code": s.get("class_code"),
                "class_name": s.get("class_name"),
                "period_num": s.get("period_num"),
                "shift": s.get("shift"),
                "subject_name": s.get("subject_name"),
                "room_number": s.get("room_number")
            })

    warning_msg = None
    if non_teaching_dates:
        warning_msg = f"សម្គាល់៖ ថ្ងៃ {', '.join(non_teaching_dates)} គ្មានម៉ោងបង្រៀនទេ (ច្បាប់ត្រូវបានគិតតែថ្ងៃដែលមានម៉ោងបង្រៀន)"

    return True, warning_msg, all_affected_slots


# ==========================================
# TEACHER ACCOUNTABILITY & DASHBOARD AGGREGATIONS
# ==========================================
def calculate_teacher_attendance_from_slots(date_str=None):
    """
    ស្រង់វត្តមានគ្រូបង្រៀនដោយស្វ័យប្រវត្តិ ផ្អែកលើការស្រង់វត្តមានសិស្សតាមម៉ោងសិក្សា៖
    - ប្រសិនបើគ្រូបង្រៀនបានស្រង់វត្តមានសិស្សសម្រាប់ម៉ោងនោះ => វត្តមាន (PRESENT)
    - ប្រសិនបើគ្រូបង្រៀនខកខានមិនបានស្រង់វត្តមានសិស្ស => អវត្តមាន (ABSENT)
    - ប្រសិនបើមានច្បាប់អនុញ្ញាតត្រឹមត្រូវ => ច្បាប់ (PERMISSION)
    """
    if not date_str:
        date_str = datetime.now().strftime("%Y-%m-%d")

    try:
        dt = datetime.strptime(date_str, "%Y-%m-%d")
        weekday_idx = dt.weekday()
    except Exception:
        weekday_idx = 0

    day_map = {0: 'ច', 1: 'អ', 2: 'ព', 3: 'ព្រ', 4: 'សុ', 5: 'ស'}
    if weekday_idx not in day_map:
        return {"date": date_str, "day_code": "អា", "total_slots": 0, "present": 0, "absent": 0, "permission": 0, "unmarked_slots": []}

    day_code = day_map[weekday_idx]
    conn = get_db_connection()
    cursor = conn.cursor()

    # Get all scheduled slots for this day where an active teacher is assigned
    slots = conn.execute("""
        SELECT ts.*, COALESCE(c.id, ts.class_id) as resolved_class_id,
               t.full_name_kh as teacher_full_name, t.teacher_code as t_code
        FROM timetable_slots ts
        LEFT JOIN classes c ON ts.class_id = c.id OR c.class_name LIKE '%' || ts.class_code || '%'
        JOIN teachers t ON ts.teacher_id = t.id
        WHERE ts.day_code = ? AND ts.teacher_id IS NOT NULL AND t.status = 'Active'
        ORDER BY ts.shift ASC, ts.period_num ASC, ts.class_code ASC
    """, (day_code,)).fetchall()

    # Get approved teacher leave requests on this date
    leaves = conn.execute("""
        SELECT person_id, reason FROM leave_requests
        WHERE person_type = 'TEACHER' AND status = 'Approved'
          AND start_date <= ? AND end_date >= ?
    """, (date_str, date_str)).fetchall()
    leave_map = {r["person_id"]: r["reason"] for r in leaves}

    total_slots = len(slots)
    present_count = 0
    absent_count = 0
    permission_count = 0
    unmarked_slots = []

    for s in slots:
        t_id = s["teacher_id"]
        c_id = s["resolved_class_id"]
        p_num = s["period_num"]
        shift = s["shift"] or ("Morning" if p_num <= 4 else "Afternoon")
        period_label = f"ម៉ោងទី {p_num}"
        session_label = f"Session {p_num}"
        class_code = s["class_code"]
        teacher_name = s["teacher_full_name"] or s["teacher_name"] or "គ្រូបង្រៀន"
        subj = s["subject_name"] or s["subject_code"] or "មុខវិជ្ជា"

        # Check if student attendance is recorded for this class and period on date_str
        # (Checks both student_attendance for absences and attendance_audit_logs for completed sessions)
        att_row = conn.execute("""
            SELECT 
                (SELECT COUNT(*) FROM student_attendance 
                 WHERE date = ? AND class_id = ? AND (period = ? OR period = ? OR period = ? OR period = 'Daily'))
                +
                (SELECT COUNT(*) FROM attendance_audit_logs 
                 WHERE date = ? AND class_id = ? AND (period = ? OR period = ? OR period = ? OR period = 'Daily'))
                as cnt
        """, (date_str, c_id, period_label, session_label, str(p_num),
              date_str, c_id, period_label, session_label, str(p_num))).fetchone()

        is_marked = bool(att_row and att_row["cnt"] and att_row["cnt"] > 0)

        if is_marked:
            status = 'PRESENT'
            reason = f"បានស្រង់វត្តមានសិស្ស {period_label} ថ្នាក់ {class_code}"
            present_count += 1
        elif t_id in leave_map:
            status = 'PERMISSION'
            reason = f"ច្បាប់សម្រាក៖ {leave_map[t_id]}"
            permission_count += 1
        else:
            status = 'ABSENT'
            reason = f"ខកខានមិនបានស្រង់វត្តមានសិស្ស {period_label} ថ្នាក់ {class_code}"
            absent_count += 1
            unmarked_slots.append({
                "slot_id": s["id"],
                "class_code": class_code,
                "period_num": p_num,
                "shift": shift,
                "teacher_id": t_id,
                "teacher_name": teacher_name,
                "subject": subj,
                "reason": reason
            })

        # Reconcile into teacher_attendance (only overwrite if recorded_by was system/automatic)
        existing = conn.execute("""
            SELECT id, recorded_by FROM teacher_attendance
            WHERE date = ? AND teacher_id = ? AND shift = ? AND period = ?
        """, (date_str, t_id, shift, period_label)).fetchone()

        if not existing:
            cursor.execute("""
                INSERT INTO teacher_attendance (date, teacher_id, shift, period, status, reason, recorded_by)
                VALUES (?, ?, ?, ?, ?, ?, 'System_Timetable')
            """, (date_str, t_id, shift, period_label, status, reason))
        elif existing["recorded_by"] in ('System_Timetable', 'Auto_Reconcile', None, ''):
            cursor.execute("""
                UPDATE teacher_attendance
                SET status = ?, reason = ?, recorded_by = 'System_Timetable'
                WHERE id = ?
            """, (status, reason, existing["id"]))

    conn.commit()
    conn.close()

    return {
        "date": date_str,
        "day_code": day_code,
        "total_slots": total_slots,
        "present": present_count,
        "absent": absent_count,
        "permission": permission_count,
        "unmarked_slots": unmarked_slots
    }


def get_timetable_accountability(date_str=None, shift=None):
    """
    ទាញយកតារាងផ្ទៀងផ្ទាត់កាលវិភាគបង្រៀនប្រចាំថ្ងៃ និងស្ថានភាពស្រង់វត្តមានជាក់ស្តែង
    """
    if not date_str:
        date_str = datetime.now().strftime("%Y-%m-%d")

    try:
        dt = datetime.strptime(date_str, "%Y-%m-%d")
        weekday_idx = dt.weekday()
    except Exception:
        weekday_idx = 0

    day_map = {0: 'ច', 1: 'អ', 2: 'ព', 3: 'ព្រ', 4: 'សុ', 5: 'ស'}
    if weekday_idx not in day_map:
        return []

    day_code = day_map[weekday_idx]
    conn = get_db_connection()
    params = [day_code]
    query = """
        SELECT ts.*, COALESCE(c.id, ts.class_id) as resolved_class_id,
               t.full_name_kh as teacher_full_name, t.teacher_code as t_code, t.phone as teacher_phone
        FROM timetable_slots ts
        LEFT JOIN classes c ON ts.class_id = c.id OR c.class_name LIKE '%' || ts.class_code || '%'
        JOIN teachers t ON ts.teacher_id = t.id
        WHERE ts.day_code = ? AND ts.teacher_id IS NOT NULL
    """
    if shift and shift != 'All':
        query += " AND ts.shift = ?"
        params.append(shift)

    query += " ORDER BY ts.shift ASC, ts.period_num ASC, ts.class_code ASC"
    slots = conn.execute(query, params).fetchall()

    leaves = conn.execute("""
        SELECT person_id, reason FROM leave_requests
        WHERE person_type = 'TEACHER' AND status = 'Approved'
          AND start_date <= ? AND end_date >= ?
    """, (date_str, date_str)).fetchall()
    leave_map = {r["person_id"]: r["reason"] for r in leaves}

    result = []
    for s in slots:
        c_id = s["resolved_class_id"]
        p_num = s["period_num"]
        p_label = f"ម៉ោងទី {p_num}"
        s_label = f"Session {p_num}"
        t_id = s["teacher_id"]

        # Check audit logs (marks session even with 0 absences)
        audit_row = conn.execute("""
            SELECT COUNT(*) as c FROM attendance_audit_logs
            WHERE date = ? AND class_id = ?
              AND (period = ? OR period = ? OR period = ? OR period = 'Daily')
        """, (date_str, c_id, p_label, s_label, str(p_num))).fetchone()
        audit_cnt = audit_row["c"] if audit_row else 0

        cnt_row = conn.execute("""
            SELECT COUNT(*) as cnt,
                   SUM(CASE WHEN status = 'PRESENT' THEN 1 ELSE 0 END) as legacy_present_cnt,
                   SUM(CASE WHEN status = 'ABSENT' THEN 1 ELSE 0 END) as absent_cnt,
                   SUM(CASE WHEN status = 'PERMISSION' THEN 1 ELSE 0 END) as permission_cnt,
                   SUM(CASE WHEN status = 'LATE' THEN 1 ELSE 0 END) as late_cnt
            FROM student_attendance
            WHERE date = ? AND class_id = ?
              AND (period = ? OR period = ? OR period = ? OR period = 'Daily')
        """, (date_str, c_id, p_label, s_label, str(p_num))).fetchone()

        is_marked = bool((cnt_row and cnt_row["cnt"] > 0) or audit_cnt > 0)

        # Class size to compute present count
        class_size_row = conn.execute("SELECT COUNT(*) as c FROM students WHERE class_id = ? AND status = 'Active'", (c_id,)).fetchone()
        class_size = class_size_row["c"] if class_size_row else 0

        absent_cnt = (cnt_row["absent_cnt"] or 0) if cnt_row else 0
        permission_cnt = (cnt_row["permission_cnt"] or 0) if cnt_row else 0
        late_cnt = (cnt_row["late_cnt"] or 0) if cnt_row else 0
        legacy_present = (cnt_row["legacy_present_cnt"] or 0) if cnt_row else 0

        if legacy_present > 0:
            present_cnt = legacy_present
        elif is_marked:
            present_cnt = max(0, class_size - (absent_cnt + permission_cnt + late_cnt))
        else:
            present_cnt = 0

        if is_marked:
            teacher_status = 'PRESENT'
            status_reason = f"បានស្រង់វត្តមានសិស្ស ({present_cnt} វត្តមាន, {absent_cnt} អវត្តមាន)"
        elif t_id in leave_map:
            teacher_status = 'PERMISSION'
            status_reason = f"ច្បាប់សម្រាក៖ {leave_map[t_id]}"
        else:
            teacher_status = 'ABSENT'
            status_reason = "ខកខានមិនបានស្រង់វត្តមានសិស្សតាមម៉ោងបង្រៀន"

        slot_dict = dict(s)
        slot_dict["is_marked"] = is_marked
        slot_dict["student_attendance_count"] = (present_cnt + absent_cnt + permission_cnt + late_cnt) if is_marked else 0
        slot_dict["teacher_status"] = teacher_status
        slot_dict["status_reason"] = status_reason
        result.append(slot_dict)

    conn.close()
    return result


def get_dashboard_stats(date_str=None):
    """ទិន្នន័យសង្ខេបសម្រាប់ Dashboard"""
    if not date_str:
        date_str = datetime.now().strftime("%Y-%m-%d")

    # 1. First automatically reconcile teacher attendance from timetable slots!
    accountability_info = calculate_teacher_attendance_from_slots(date_str)

    conn = get_db_connection()

    total_teachers = conn.execute("SELECT COUNT(*) FROM teachers WHERE status = 'Active'").fetchone()[0]
    total_students = conn.execute("SELECT COUNT(*) FROM students WHERE status = 'Active'").fetchone()[0]
    total_classes = conn.execute("SELECT COUNT(*) FROM classes").fetchone()[0]

    # Teacher today stats
    teacher_stats = conn.execute("""
        SELECT 
            SUM(CASE WHEN status = 'PRESENT' THEN 1 ELSE 0 END) as present,
            SUM(CASE WHEN status = 'PERMISSION' THEN 1 ELSE 0 END) as permission,
            SUM(CASE WHEN status = 'ABSENT' THEN 1 ELSE 0 END) as absent,
            SUM(CASE WHEN status = 'LATE' THEN 1 ELSE 0 END) as late,
            COUNT(*) as total_recorded
        FROM teacher_attendance
        WHERE date = ?
    """, (date_str,)).fetchone()

    # Student today stats
    student_stats = conn.execute("""
        SELECT 
            SUM(CASE WHEN status = 'PRESENT' THEN 1 ELSE 0 END) as present,
            SUM(CASE WHEN status = 'PERMISSION' THEN 1 ELSE 0 END) as permission,
            SUM(CASE WHEN status = 'ABSENT' THEN 1 ELSE 0 END) as absent,
            SUM(CASE WHEN status = 'LATE' THEN 1 ELSE 0 END) as late,
            COUNT(*) as total_recorded
        FROM student_attendance
        WHERE date = ?
    """, (date_str,)).fetchone()

    st_perm = (student_stats["permission"] or 0) if student_stats else 0
    st_ab = (student_stats["absent"] or 0) if student_stats else 0
    st_late = (student_stats["late"] or 0) if student_stats else 0
    legacy_present = (student_stats["present"] or 0) if student_stats else 0

    if legacy_present > 0:
        st_present = legacy_present
    else:
        # Sparse mode: compute present count from classes recorded today
        recorded_classes = conn.execute("""
            SELECT al.class_id, COUNT(DISTINCT al.period) as sessions_cnt,
                   (SELECT COUNT(*) FROM students WHERE class_id = al.class_id AND status = 'Active') as class_size
            FROM attendance_audit_logs al
            WHERE al.date = ?
            GROUP BY al.class_id
        """, (date_str,)).fetchall()
        total_student_sessions = sum((r["sessions_cnt"] or 0) * (r["class_size"] or 0) for r in recorded_classes)
        st_present = max(0, total_student_sessions - (st_perm + st_ab + st_late))

    # Absent teachers today
    absent_teachers = conn.execute("""
        SELECT ta.*, t.full_name_kh, t.subject, t.phone, sub.full_name_kh as substitute_name
        FROM teacher_attendance ta
        JOIN teachers t ON ta.teacher_id = t.id
        LEFT JOIN teachers sub ON ta.substitute_teacher_id = sub.id
        WHERE ta.date = ? AND ta.status IN ('ABSENT', 'PERMISSION', 'LATE')
        ORDER BY ta.shift ASC, t.full_name_kh ASC
    """, (date_str,)).fetchall()

    # Absent students today
    absent_students = conn.execute("""
        SELECT sa.*, s.full_name_kh, s.student_code, s.gender, s.parent_phone, c.class_name
        FROM student_attendance sa
        JOIN students s ON sa.student_id = s.id
        JOIN classes c ON sa.class_id = c.id
        WHERE sa.date = ? AND sa.status IN ('ABSENT', 'PERMISSION', 'LATE')
        ORDER BY c.class_name ASC, s.full_name_kh ASC
    """, (date_str,)).fetchall()

    conn.close()

    return {
        "date": date_str,
        "total_teachers": total_teachers,
        "total_students": total_students,
        "total_classes": total_classes,
        "timetable_reconcile": accountability_info,
        "unmarked_slots": accountability_info.get("unmarked_slots", []),
        "unmarked_slots_count": len(accountability_info.get("unmarked_slots", [])),
        "teachers": {
            "present": teacher_stats["present"] or 0,
            "permission": teacher_stats["permission"] or 0,
            "absent": teacher_stats["absent"] or 0,
            "late": teacher_stats["late"] or 0,
            "total_recorded": teacher_stats["total_recorded"] or 0,
            "absent_list": [dict(r) for r in absent_teachers]
        },
        "students": {
            "present": st_present,
            "permission": st_perm,
            "absent": st_ab,
            "late": st_late,
            "total_recorded": st_present + st_perm + st_ab + st_late,
            "absent_list": [dict(r) for r in absent_students]
        }
    }


def get_student_report_data(class_id=None, start_date=None, end_date=None):
    """របាយការណ៍អវត្តមានសិស្សតាមចន្លោះកាលបរិច្ឆេទ"""
    conn = get_db_connection()
    params = []
    where_clauses = ["1=1"]

    if class_id:
        where_clauses.append("s.class_id = ?")
        params.append(class_id)
    if start_date:
        where_clauses.append("sa.date >= ?")
        params.append(start_date)
    if end_date:
        where_clauses.append("sa.date <= ?")
        params.append(end_date)

    where_sql = " AND ".join(where_clauses)

    query = f"""
        SELECT s.id, s.student_code, s.full_name_kh, s.gender, c.class_name,
               COUNT(sa.id) as total_sessions,
               SUM(CASE WHEN sa.status = 'PRESENT' THEN 1 ELSE 0 END) as present_count,
               SUM(CASE WHEN sa.status = 'PERMISSION' THEN 1 ELSE 0 END) as permission_count,
               SUM(CASE WHEN sa.status = 'ABSENT' THEN 1 ELSE 0 END) as absent_count,
               SUM(CASE WHEN sa.status = 'LATE' THEN 1 ELSE 0 END) as late_count
        FROM students s
        JOIN classes c ON s.class_id = c.id
        LEFT JOIN student_attendance sa ON s.id = sa.student_id
        WHERE {where_sql}
        GROUP BY s.id
        ORDER BY c.class_name ASC, s.full_name_kh ASC
    """
    rows = conn.execute(query, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_teacher_report_data(start_date=None, end_date=None):
    """របាយការណ៍អវត្តមានគ្រូតាមចន្លោះកាលបរិច្ឆេទ"""
    conn = get_db_connection()
    params = []
    where_clauses = ["1=1"]

    if start_date:
        where_clauses.append("ta.date >= ?")
        params.append(start_date)
    if end_date:
        where_clauses.append("ta.date <= ?")
        params.append(end_date)

    where_sql = " AND ".join(where_clauses)

    query = f"""
        SELECT t.id, t.teacher_code, t.full_name_kh, t.gender, t.subject,
               COUNT(ta.id) as total_sessions,
               SUM(CASE WHEN ta.status = 'PRESENT' THEN 1 ELSE 0 END) as present_count,
               SUM(CASE WHEN ta.status = 'PERMISSION' THEN 1 ELSE 0 END) as permission_count,
               SUM(CASE WHEN ta.status = 'ABSENT' THEN 1 ELSE 0 END) as absent_count,
               SUM(CASE WHEN ta.status = 'LATE' THEN 1 ELSE 0 END) as late_count
        FROM teachers t
        LEFT JOIN teacher_attendance ta ON t.id = ta.teacher_id
        WHERE {where_sql}
        GROUP BY t.id
        ORDER BY t.full_name_kh ASC
    """
    rows = conn.execute(query, params).fetchall()
    conn.close()
    return [dict(r) for r in rows]


# ==========================================
# SETTINGS OPERATIONS
# ==========================================
def get_setting(key, default=None):
    conn = get_db_connection()
    row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    conn.close()
    return row["value"] if row else default


def set_setting(key, value):
    conn = get_db_connection()
    conn.execute("""
        INSERT INTO settings (key, value, updated_at)
        VALUES (?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = CURRENT_TIMESTAMP
    """, (key, str(value)))
    conn.commit()
    conn.close()


def get_all_settings():
    conn = get_db_connection()
    rows = conn.execute("SELECT key, value FROM settings").fetchall()
    conn.close()
    return {r["key"]: r["value"] for r in rows}


# ==========================================
# TIMETABLE OPERATIONS
# ==========================================
def get_timetable_by_class(class_code):
    """ទាញយកកាលវិភាគតាមថ្នាក់រៀន (ចន្ទ ដល់ សៅរ៍)"""
    conn = get_db_connection()
    rows = conn.execute("""
        SELECT ts.*, t.full_name_kh as teacher_full_name, t.phone as teacher_phone
        FROM timetable_slots ts
        LEFT JOIN teachers t ON ts.teacher_id = t.id
        WHERE ts.class_code = ?
        ORDER BY 
            CASE ts.day_code
                WHEN 'ច' THEN 1
                WHEN 'អ' THEN 2
                WHEN 'ព' THEN 3
                WHEN 'ព្រ' THEN 4
                WHEN 'សុ' THEN 5
                WHEN 'ស' THEN 6
                ELSE 7
            END,
            ts.period_num ASC
    """, (class_code,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_timetable_by_teacher(teacher_id):
    """ទាញយកកាលវិភាគតាមគ្រូបង្រៀន"""
    conn = get_db_connection()
    rows = conn.execute("""
        SELECT ts.*, c.class_name
        FROM timetable_slots ts
        LEFT JOIN classes c ON ts.class_id = c.id
        WHERE ts.teacher_id = ?
        ORDER BY 
            CASE ts.day_code
                WHEN 'ច' THEN 1
                WHEN 'អ' THEN 2
                WHEN 'ព' THEN 3
                WHEN 'ព្រ' THEN 4
                WHEN 'សុ' THEN 5
                WHEN 'ស' THEN 6
                ELSE 7
            END,
            ts.period_num ASC
    """, (teacher_id,)).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_all_timetable_classes():
    """ទាញយកបញ្ជីថ្នាក់រៀនដែលមានក្នុងកាលវិភាគ"""
    conn = get_db_connection()
    rows = conn.execute("""
        SELECT DISTINCT class_code FROM timetable_slots ORDER BY class_code ASC
    """).fetchall()
    conn.close()
    return [r["class_code"] for r in rows]


def get_class_by_code(class_code):
    """ទាញយកព័ត៌មានថ្នាក់រៀនតាមកូដ (ឧ. 7A)"""
    conn = get_db_connection()
    c = conn.execute("""
        SELECT c.*, t.full_name_kh as homeroom_teacher_name
        FROM classes c
        LEFT JOIN teachers t ON c.homeroom_teacher_id = t.id
        WHERE c.class_name LIKE ? OR c.class_name = ?
        LIMIT 1
    """, (f"%{class_code}%", class_code)).fetchone()
    conn.close()
    return dict(c) if c else None


def get_teacher_today_slots(teacher_id, day_code, today_date=None, current_dt=None):
    """
    ទាញយកកាលវិភាគបង្រៀនរបស់គ្រូនៅថ្ងៃនេះ រួមទាំងស្ថានភាពស្រង់វត្តមានសិស្ស និងភាពអាចស្រង់បាន (Timing & Disabled status)
    - ថ្នាក់ដែលត្រូវម៉ោង (Active period): can_take_attendance=True (ឬ can_edit_attendance=True ក្នុង ៣០ នាទីដំបូង)
    - ថ្នាក់ដែលបានស្រង់រួច (ក្រៅពីម៉ោងសកម្ម): Disabled (ចាក់សោ)
    - ថ្នាក់ដែលមិនទាន់ដល់ម៉ោង (Upcoming) ឬផុតម៉ោង (Past): Disabled (មិនទាន់ដល់ម៉ោង / ផុតម៉ោងស្រង់)
    """
    if not today_date:
        today_date = datetime.now().strftime("%Y-%m-%d")
    if current_dt is None:
        current_dt = datetime.now()

    conn = get_db_connection()
    rows = conn.execute("""
        SELECT ts.*, COALESCE(c.id, ts.class_id) as class_db_id, c.class_name,
               (SELECT COUNT(*) FROM students s WHERE s.class_id = COALESCE(c.id, ts.class_id) AND s.status = 'Active') as student_count,
               (SELECT COUNT(*) FROM student_attendance sa 
                WHERE sa.date = ? 
                  AND sa.class_id = COALESCE(c.id, ts.class_id)
                  AND (sa.period = 'ម៉ោងទី ' || ts.period_num OR sa.period = 'Session ' || ts.period_num OR sa.period = CAST(ts.period_num AS TEXT) OR sa.period = 'Daily')
               ) as attendance_marked_count
        FROM timetable_slots ts
        LEFT JOIN classes c ON ts.class_id = c.id OR c.class_name LIKE '%' || ts.class_code || '%'
        WHERE ts.teacher_id = ? AND ts.day_code = ?
        ORDER BY ts.period_num ASC
    """, (today_date, teacher_id, day_code)).fetchall()
    conn.close()

    # Period schedule mapping indexed by db_period_num
    period_map = {p["db_period_num"]: p for p in PERIOD_SCHEDULE}
    today_weekday_code = DAY_MAP.get(current_dt.weekday(), {}).get("code", "")
    is_same_day = (day_code == today_weekday_code and today_date == current_dt.strftime("%Y-%m-%d"))
    cur_minutes = current_dt.hour * 60 + current_dt.minute

    result = []
    for r in rows:
        d = dict(r)
        p_info = period_map.get(d["period_num"], {})
        d["time_range"] = p_info.get("time_range", "")
        d["period_label"] = p_info.get("label", f"ម៉ោងទី {d['period_num']}")
        d["short_label"] = p_info.get("short_label", f"ម៉ោងទី {d['period_num']}")
        d["is_marked"] = bool(d.get("attendance_marked_count", 0) and d["attendance_marked_count"] > 0)

        start_min = (p_info.get("start_hour", 0) * 60) + p_info.get("start_min", 0)
        end_min = (p_info.get("end_hour", 0) * 60) + p_info.get("end_min", 0)

        if not is_same_day:
            d["timing_status"] = "OTHER_DAY"
            d["is_current_active"] = False
            d["can_take_attendance"] = False
            d["can_edit_attendance"] = False
            d["disabled_reason"] = "មិនមែនថ្ងៃនេះទេ"
        elif start_min <= cur_minutes < end_min:
            d["timing_status"] = "ACTIVE"
            d["is_current_active"] = True
            if not d["is_marked"]:
                d["can_take_attendance"] = True
                d["can_edit_attendance"] = False
                d["disabled_reason"] = ""
            else:
                d["can_take_attendance"] = False
                # If within first 30 minutes, can edit
                if cur_minutes < (start_min + 30):
                    d["can_edit_attendance"] = True
                    d["disabled_reason"] = ""
                else:
                    d["can_edit_attendance"] = False
                    d["disabled_reason"] = "ផុតម៉ោងអនុញ្ញាតកែប្រែ (បានបញ្ចូលរួច)"
        elif cur_minutes < start_min:
            d["timing_status"] = "UPCOMING"
            d["is_current_active"] = False
            d["can_take_attendance"] = False
            d["can_edit_attendance"] = False
            d["disabled_reason"] = "មិនទាន់ដល់ម៉ោងបង្រៀន"
        else:
            d["timing_status"] = "PAST"
            d["is_current_active"] = False
            d["can_take_attendance"] = False
            d["can_edit_attendance"] = False
            if d["is_marked"]:
                d["disabled_reason"] = "បានស្រង់រួច (បិទបញ្ចប់)"
            else:
                d["disabled_reason"] = "ផុតម៉ោងស្រង់វត្តមាន"

        result.append(d)
    return result


def get_all_subjects():
    """ទាញយកមុខវិជ្ជាទាំងអស់ពី database"""
    conn = get_db_connection()
    try:
        rows = conn.execute("SELECT * FROM subjects ORDER BY name_kh ASC").fetchall()
        if rows:
            conn.close()
            return [dict(r) for r in rows]
    except Exception:
        pass

    rows = conn.execute("""
        SELECT DISTINCT subject_name as name_kh, subject_code as code
        FROM timetable_slots
        WHERE subject_name IS NOT NULL AND subject_name != ''
        ORDER BY subject_name ASC
    """).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_timetable_slot_by_id(slot_id):
    """ទាញយកម៉ោងបង្រៀនមួយតាម ID"""
    conn = get_db_connection()
    row = conn.execute("""
        SELECT ts.*, t.full_name_kh as teacher_full_name, c.class_name
        FROM timetable_slots ts
        LEFT JOIN teachers t ON ts.teacher_id = t.id
        LEFT JOIN classes c ON ts.class_id = c.id
        WHERE ts.id = ?
    """, (slot_id,)).fetchone()
    conn.close()
    return dict(row) if row else None


def save_timetable_slot(slot_id=None, class_code="", teacher_id=None, day_code="", period_num=1, shift=None, subject_name="", room_number=""):
    """
    បង្កើត ឬកែប្រែកាលវិភាគបង្រៀន (Timetable Slot)
    """
    conn = get_db_connection()
    cursor = conn.cursor()

    day_mapping = {
        'ច': 'ចន្ទ',
        'អ': 'អង្គារ',
        'ព': 'ពុធ',
        'ព្រ': 'ព្រហស្បតិ៍',
        'សុ': 'សុក្រ',
        'ស': 'សៅរ៍'
    }
    day_name = day_mapping.get(day_code, "ចន្ទ")

    period_num = int(period_num)
    if not shift:
        shift = "Morning" if period_num <= 4 else "Afternoon"

    # Resolve class_id
    class_row = conn.execute("SELECT id, room_number FROM classes WHERE class_name LIKE ? OR class_name = ? LIMIT 1",
                             (f"%{class_code}%", class_code)).fetchone()
    class_id = class_row["id"] if class_row else None
    if not room_number and class_row and class_row["room_number"]:
        room_number = class_row["room_number"]

    # Resolve teacher info
    teacher_code = ""
    teacher_name = ""
    if teacher_id:
        teacher_row = conn.execute("SELECT id, teacher_code, full_name_kh FROM teachers WHERE id = ?", (teacher_id,)).fetchone()
        if teacher_row:
            teacher_code = teacher_row["teacher_code"]
            teacher_name = teacher_row["full_name_kh"]

    # Resolve subject_code
    subject_code = ""
    if subject_name:
        sub_row = conn.execute("SELECT code FROM subjects WHERE name_kh = ? LIMIT 1", (subject_name,)).fetchone()
        if sub_row:
            subject_code = sub_row["code"]

    if slot_id:
        cursor.execute("""
            UPDATE timetable_slots
            SET class_id = ?, class_code = ?, day_code = ?, day_name = ?, period_num = ?, shift = ?,
                subject_code = ?, subject_name = ?, teacher_id = ?, teacher_code = ?, teacher_name = ?, room_number = ?
            WHERE id = ?
        """, (class_id, class_code, day_code, day_name, period_num, shift, subject_code, subject_name, teacher_id, teacher_code, teacher_name, room_number, slot_id))
        saved_id = slot_id
    else:
        cursor.execute("""
            INSERT INTO timetable_slots (class_id, class_code, day_code, day_name, period_num, shift, subject_code, subject_name, teacher_id, teacher_code, teacher_name, room_number)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(class_code, day_code, period_num) DO UPDATE SET
                class_id = excluded.class_id,
                day_name = excluded.day_name,
                shift = excluded.shift,
                subject_code = excluded.subject_code,
                subject_name = excluded.subject_name,
                teacher_id = excluded.teacher_id,
                teacher_code = excluded.teacher_code,
                teacher_name = excluded.teacher_name,
                room_number = excluded.room_number
        """, (class_id, class_code, day_code, day_name, period_num, shift, subject_code, subject_name, teacher_id, teacher_code, teacher_name, room_number))
        saved_id = cursor.lastrowid

    conn.commit()
    conn.close()
    return saved_id


def delete_timetable_slot(slot_id):
    """លុបម៉ោងបង្រៀនចេញពីកាលវិភាគ"""
    conn = get_db_connection()
    conn.execute("DELETE FROM timetable_slots WHERE id = ?", (slot_id,))
    conn.commit()
    conn.close()
    return True


def import_timetable_from_excel(file_path, mode="merge"):
    """
    Import master school timetable from an Excel file into timetable_slots.
    mode: 'replace' (clear all slots first) or 'merge' (upsert slots).
    """
    import openpyxl

    wb = openpyxl.load_workbook(file_path, data_only=True)
    if "កាលវិភាគរួម" in wb.sheetnames:
        ws = wb["កាលវិភាគរួម"]
    else:
        ws = wb.active

    day_clean_map = {
        "ច": "ច", "ចន្ទ": "ច", "mon": "ច", "monday": "ច",
        "អ": "អ", "អង្គារ": "អ", "tue": "អ", "tuesday": "អ",
        "ព": "ព", "ពុធ": "ព", "wed": "ព", "wednesday": "ព",
        "ព្រ": "ព្រ", "ព្រហ": "ព្រ", "ព្រហស្បតិ៍": "ព្រ", "thu": "ព្រ", "thursday": "ព្រ",
        "សុ": "សុ", "សុក្រ": "សុ", "fri": "សុ", "friday": "សុ",
        "ស": "ស", "សៅរ៍": "ស", "sat": "ស", "saturday": "ស"
    }

    conn = get_db_connection()
    all_teachers = conn.execute("SELECT id, teacher_code, full_name_kh, subject FROM teachers").fetchall()
    legacy_slots = conn.execute("SELECT DISTINCT teacher_code, teacher_id FROM timetable_slots WHERE teacher_id IS NOT NULL AND teacher_code != ''").fetchall()
    conn.close()

    teacher_by_id = {t["id"]: dict(t) for t in all_teachers}
    teacher_by_code = {str(t["teacher_code"]).strip().lower(): dict(t) for t in all_teachers}
    teacher_by_name = {str(t["full_name_kh"]).strip(): dict(t) for t in all_teachers}

    # Also map legacy timetable codes (e.g. H3, M1, K3) to the teacher object
    for lr in legacy_slots:
        code_str = str(lr["teacher_code"]).strip().lower()
        if code_str not in teacher_by_code and lr["teacher_id"] in teacher_by_id:
            teacher_by_code[code_str] = teacher_by_id[lr["teacher_id"]]

    header_row_idx = 1
    col_map = {}
    for r in range(1, min(15, ws.max_row + 1)):
        row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
        for c_idx, val in enumerate(row_vals, start=1):
            val_lower = val.lower()
            if "ថ្នាក់" in val or "class" in val_lower:
                col_map["class_code"] = c_idx
            elif "ថ្ងៃ" in val or "day" in val_lower:
                col_map["day_code"] = c_idx
            elif "ម៉ោង" in val or "period" in val_lower:
                col_map["period_num"] = c_idx
            elif "មុខវិជ្ជា" in val or "subject" in val_lower:
                col_map["subject_name"] = c_idx
            elif "គ្រូ" in val or "teacher" in val_lower:
                col_map["teacher"] = c_idx
            elif "បន្ទប់" in val or "room" in val_lower:
                col_map["room_number"] = c_idx
        if len(col_map) >= 4:
            header_row_idx = r
            break

    if len(col_map) < 4:
        col_map = {
            "class_code": 2,
            "day_code": 3,
            "period_num": 4,
            "subject_name": 5,
            "teacher": 6,
            "room_number": 7
        }
        header_row_idx = 5

    if mode == "replace":
        conn = get_db_connection()
        conn.execute("DELETE FROM timetable_slots")
        conn.commit()
        conn.close()

    imported_count = 0
    errors = []

    for r in range(header_row_idx + 1, ws.max_row + 1):
        raw_class = str(ws.cell(row=r, column=col_map.get("class_code", 2)).value or "").strip()
        raw_day = str(ws.cell(row=r, column=col_map.get("day_code", 3)).value or "").strip()
        raw_period = str(ws.cell(row=r, column=col_map.get("period_num", 4)).value or "").strip()
        raw_subject = str(ws.cell(row=r, column=col_map.get("subject_name", 5)).value or "").strip()
        raw_teacher = str(ws.cell(row=r, column=col_map.get("teacher", 6)).value or "").strip()
        raw_room = str(ws.cell(row=r, column=col_map.get("room_number", 7)).value or "").strip()

        if not raw_class and not raw_day and not raw_period and not raw_subject:
            continue

        if not raw_class or not raw_day or not raw_period or not raw_teacher:
            errors.append(f"ជួរដេកទី {r}៖ ខ្វះព័ត៌មានចាំបាច់ (ថ្នាក់, ថ្ងៃ, ម៉ោង, ឬ គ្រូបង្រៀន)")
            continue

        class_code = raw_class.replace("ថ្នាក់ទី ", "").replace("ថ្នាក់ ", "").strip()

        cleaned_day = raw_day.strip().lower()
        day_code = day_clean_map.get(cleaned_day, raw_day.strip())
        if day_code not in ["ច", "អ", "ព", "ព្រ", "សុ", "ស"]:
            errors.append(f"ជួរដេកទី {r}៖ កូដថ្ងៃ '{raw_day}' មិនត្រឹមត្រូវទេ (ត្រូវជា ច, អ, ព, ព្រ, សុ, ស)")
            continue

        period_num_digits = ''.join(filter(str.isdigit, raw_period))
        if not period_num_digits:
            errors.append(f"ជួរដេកទី {r}៖ លេខម៉ោង '{raw_period}' មិនត្រឹមត្រូវទេ")
            continue
        period_num = int(period_num_digits)
        if period_num < 1 or period_num > 8:
            errors.append(f"ជួរដេកទី {r}៖ ម៉ោងទី {period_num} ត្រូវតែស្ថិតក្នុងចន្លោះពី ១ ដល់ ៨")
            continue

        matched_teacher = teacher_by_code.get(raw_teacher.lower())
        if not matched_teacher:
            matched_teacher = teacher_by_name.get(raw_teacher)
        if not matched_teacher:
            # Check combined format e.g. "2890800247 - សួរ ចន្ទ្រា" or "សួរ ចន្ទ្រា (2890800247)"
            tokens = re.split(r'[\s\-\(\)\:\/]+', raw_teacher)
            for tok in tokens:
                tok_clean = tok.strip()
                if not tok_clean:
                    continue
                if tok_clean.lower() in teacher_by_code:
                    matched_teacher = teacher_by_code[tok_clean.lower()]
                    break
                elif tok_clean in teacher_by_name:
                    matched_teacher = teacher_by_name[tok_clean]
                    break
        if not matched_teacher:
            for t_name, t_obj in teacher_by_name.items():
                if raw_teacher in t_name or t_name in raw_teacher:
                    matched_teacher = t_obj
                    break

        if not matched_teacher:
            errors.append(f"ជួរដេកទី {r}៖ រកមិនឃើញគ្រូបង្រៀន '{raw_teacher}' ក្នុងបញ្ជីគ្រូឡើយ")
            continue

        teacher_id = matched_teacher["id"]
        subject_name = raw_subject or matched_teacher.get("subject", "")
        shift = "Morning" if period_num <= 4 else "Afternoon"

        save_timetable_slot(
            class_code=class_code,
            teacher_id=teacher_id,
            day_code=day_code,
            period_num=period_num,
            shift=shift,
            subject_name=subject_name,
            room_number=raw_room
        )
        imported_count += 1

    return {
        "success": True,
        "imported_count": imported_count,
        "error_count": len(errors),
        "errors": errors[:30]
    }


# ==========================================
# USER AUTHENTICATION & MANAGEMENT
# ==========================================
def create_user(username, password, role="teacher", full_name_kh="", teacher_id=None, phone="", plain_hint=None):
    """បង្កើតគណនី User ថ្មីជាមួយ Hashed Password"""
    conn = get_db_connection()
    cursor = conn.cursor()
    pwd_hash = generate_password_hash(password)
    hint = plain_hint if plain_hint is not None else password

    cursor.execute("""
        INSERT INTO users (username, password_hash, plain_password_hint, role, full_name_kh, teacher_id, phone)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(username) DO UPDATE SET
            password_hash = excluded.password_hash,
            plain_password_hint = excluded.plain_password_hint,
            role = excluded.role,
            full_name_kh = excluded.full_name_kh,
            teacher_id = excluded.teacher_id,
            phone = excluded.phone
    """, (username, pwd_hash, hint, role, full_name_kh, teacher_id, phone))
    conn.commit()
    uid = cursor.lastrowid
    conn.close()
    return uid


def authenticate_user(username, password):
    """
    ផ្ទៀងផ្ទាត់ Username ឬ លេខទូរស័ព្ទ និង Password
    ត្រឡប់ dict នៃ User ប្រសិនបើត្រឹមត្រូវ, None ប្រសិនបើខុស
    """
    if not username or not password:
        return None

    clean_username = str(username).strip()
    phone_clean = clean_username.replace(' ', '').replace('-', '')

    conn = get_db_connection()
    user = conn.execute("""
        SELECT * FROM users 
        WHERE (username = ? OR phone = ? OR REPLACE(REPLACE(COALESCE(phone, ''), ' ', ''), '-', '') = ?) AND is_active = 1
        LIMIT 1
    """, (clean_username, clean_username, phone_clean)).fetchone()
    if not user:
        conn.close()
        return None

    if check_password_hash(user["password_hash"], password):
        # Update last login
        conn.execute("UPDATE users SET last_login = CURRENT_TIMESTAMP WHERE id = ?", (user["id"],))
        conn.commit()
        user_dict = dict(user)
        conn.close()
        return user_dict

    conn.close()
    return None


def change_user_password(user_id, new_password):
    """ផ្លាស់ប្តូរលេខសម្ងាត់ផ្ទាល់ខ្លួន"""
    conn = get_db_connection()
    pwd_hash = generate_password_hash(new_password)
    conn.execute("""
        UPDATE users 
        SET password_hash = ?, plain_password_hint = '[បានកែប្រែដោយផ្ទាល់]'
        WHERE id = ?
    """, (pwd_hash, user_id))
    conn.commit()
    conn.close()
    return True


def reset_user_password(user_id, new_password="123456"):
    """Reset លេខសម្ងាត់ដោយ Admin"""
    conn = get_db_connection()
    pwd_hash = generate_password_hash(new_password)
    conn.execute("""
        UPDATE users 
        SET password_hash = ?, plain_password_hint = ?
        WHERE id = ?
    """, (pwd_hash, new_password, user_id))
    conn.commit()
    conn.close()
    return True


def get_all_users():
    """ទាញយកបញ្ជី Users ទាំងអស់"""
    conn = get_db_connection()
    rows = conn.execute("""
        SELECT u.*, t.subject, t.teacher_code
        FROM users u
        LEFT JOIN teachers t ON u.teacher_id = t.id
        ORDER BY 
            CASE u.role WHEN 'admin' THEN 1 ELSE 2 END,
            u.full_name_kh ASC
    """).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def get_user_by_id(user_id):
    conn = get_db_connection()
    row = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
    conn.close()
    return dict(row) if row else None


def get_user_by_username(username):
    conn = get_db_connection()
    row = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
    conn.close()
    return dict(row) if row else None


def init_default_users():
    """
    បង្កើតគណនី Admin និងគណនីសម្រាប់គ្រូបង្រៀនទាំងអស់ដោយស្វ័យប្រវត្តិ
    """
    init_db()
    conn = get_db_connection()

    # 1. Ensure Admin Account
    admin_user = conn.execute("SELECT id FROM users WHERE username = 'admin'").fetchone()
    if not admin_user:
        create_user(
            username="admin",
            password="admin123",
            role="admin",
            full_name_kh="គណៈគ្រប់គ្រងសាលា (Administrator)",
            teacher_id=None,
            phone="012 888 999",
            plain_hint="admin123"
        )
        print("[OK] Created default Admin account: admin / admin123")

    # 2. Ensure Teacher Accounts for all teachers in database
    teachers = conn.execute("SELECT * FROM teachers").fetchall()
    created_count = 0

    for t in teachers:
        t_code = str(t["teacher_code"]).strip()
        t_name = t["full_name_kh"]
        t_id = t["id"]
        t_phone = t["phone"] or ""

        existing = conn.execute("SELECT id FROM users WHERE username = ?", (t_code,)).fetchone()
        if not existing:
            create_user(
                username=t_code,
                password="123456",
                role="teacher",
                full_name_kh=t_name,
                teacher_id=t_id,
                phone=t_phone,
                plain_hint="123456"
            )
            created_count += 1

    conn.close()
    print(f"[OK] Initialized user accounts. Created {created_count} new teacher accounts (Default password: 123456).")


# ==========================================
# TIME WINDOW & PERIOD SCHEDULE OPERATIONS
# ==========================================

PERIOD_SCHEDULE = [
    # Morning Shift: 07:00 - 11:00
    {
        "shift": "Morning",
        "shift_kh": "ពេលព្រឹក",
        "period_display_num": 1,
        "db_period_num": 1,
        "start_hour": 7,
        "start_min": 0,
        "end_hour": 8,
        "end_min": 0,
        "label": "ម៉ោងទី ១ (០៧:០០ - ០៨:០០)",
        "short_label": "ម៉ោងទី ១",
        "time_range": "07:00 - 08:00"
    },
    {
        "shift": "Morning",
        "shift_kh": "ពេលព្រឹក",
        "period_display_num": 2,
        "db_period_num": 2,
        "start_hour": 8,
        "start_min": 0,
        "end_hour": 9,
        "end_min": 0,
        "label": "ម៉ោងទី ២ (០៨:០០ - ០៩:០០)",
        "short_label": "ម៉ោងទី ២",
        "time_range": "08:00 - 09:00"
    },
    {
        "shift": "Morning",
        "shift_kh": "ពេលព្រឹក",
        "period_display_num": 3,
        "db_period_num": 3,
        "start_hour": 9,
        "start_min": 0,
        "end_hour": 10,
        "end_min": 0,
        "label": "ម៉ោងទី ៣ (០៩:០០ - ១០:០០)",
        "short_label": "ម៉ោងទី ៣",
        "time_range": "09:00 - 10:00"
    },
    {
        "shift": "Morning",
        "shift_kh": "ពេលព្រឹក",
        "period_display_num": 4,
        "db_period_num": 4,
        "start_hour": 10,
        "start_min": 0,
        "end_hour": 11,
        "end_min": 0,
        "label": "ម៉ោងទី ៤ (១០:០០ - ១១:០០)",
        "short_label": "ម៉ោងទី ៤",
        "time_range": "10:00 - 11:00"
    },
    # Afternoon Shift: 13:00 - 17:00
    {
        "shift": "Afternoon",
        "shift_kh": "ពេលរសៀល",
        "period_display_num": 1,
        "db_period_num": 5,
        "start_hour": 13,
        "start_min": 0,
        "end_hour": 14,
        "end_min": 0,
        "label": "ម៉ោងទី ១ (១៣:០០ - ១៤:០០)",
        "short_label": "ម៉ោងទី ១ (រសៀល)",
        "time_range": "13:00 - 14:00"
    },
    {
        "shift": "Afternoon",
        "shift_kh": "ពេលរសៀល",
        "period_display_num": 2,
        "db_period_num": 6,
        "start_hour": 14,
        "start_min": 0,
        "end_hour": 15,
        "end_min": 0,
        "label": "ម៉ោងទី ២ (១៤:០០ - ១៥:០០)",
        "short_label": "ម៉ោងទី ២ (រសៀល)",
        "time_range": "14:00 - 15:00"
    },
    {
        "shift": "Afternoon",
        "shift_kh": "ពេលរសៀល",
        "period_display_num": 3,
        "db_period_num": 7,
        "start_hour": 15,
        "start_min": 0,
        "end_hour": 16,
        "end_min": 0,
        "label": "ម៉ោងទី ៣ (១៥:០០ - ១៦:០០)",
        "short_label": "ម៉ោងទី ៣ (រសៀល)",
        "time_range": "15:00 - 16:00"
    },
    {
        "shift": "Afternoon",
        "shift_kh": "ពេលរសៀល",
        "period_display_num": 4,
        "db_period_num": 8,
        "start_hour": 16,
        "start_min": 0,
        "end_hour": 17,
        "end_min": 0,
        "label": "ម៉ោងទី ៤ (១៦:០០ - ១៧:០០)",
        "short_label": "ម៉ោងទី ៤ (រសៀល)",
        "time_range": "16:00 - 17:00"
    }
]

DAY_MAP = {
    0: {"code": "ច", "name": "ចន្ទ", "en": "Monday"},
    1: {"code": "អ", "name": "អង្គារ", "en": "Tuesday"},
    2: {"code": "ព", "name": "ពុធ", "en": "Wednesday"},
    3: {"code": "ព្រ", "name": "ព្រហស្បតិ៍", "en": "Thursday"},
    4: {"code": "សុ", "name": "សុក្រ", "en": "Friday"},
    5: {"code": "ស", "name": "សៅរ៍", "en": "Saturday"},
    6: {"code": "អា", "name": "អាទិត្យ", "en": "Sunday"}
}


def get_current_period_info(current_dt=None):
    """
    កំណត់ព័ត៌មានម៉ោងសិក្សាបច្ចុប្បន្ន (Period info, Phase, Seconds remaining)
    """
    if current_dt is None:
        current_dt = datetime.now()

    day_info = DAY_MAP.get(current_dt.weekday(), DAY_MAP[0])
    cur_minutes = current_dt.hour * 60 + current_dt.minute

    matched_period = None
    for p in PERIOD_SCHEDULE:
        start_min = p["start_hour"] * 60 + p["start_min"]
        end_min = p["end_hour"] * 60 + p["end_min"]
        if start_min <= cur_minutes < end_min:
            matched_period = p.copy()
            break

    # Calculate exact seconds until the next hour (turn of the hour :00:00)
    sec_until_next_hour = (60 - current_dt.minute) * 60 - current_dt.second
    if sec_until_next_hour <= 0:
        sec_until_next_hour = 3600

    if not matched_period:
        upcoming = None
        for p in PERIOD_SCHEDULE:
            start_min = p["start_hour"] * 60 + p["start_min"]
            if cur_minutes < start_min:
                upcoming = p
                break

        return {
            "has_active_period": False,
            "period": None,
            "day": day_info,
            "current_time_str": current_dt.strftime("%H:%M:%S"),
            "current_date_str": current_dt.strftime("%Y-%m-%d"),
            "current_date_display": current_dt.strftime("%d/%m/%Y"),
            "sec_until_next_hour": sec_until_next_hour,
            "upcoming_period": upcoming,
            "phase": "no_period",
            "phase_kh": "ក្រៅម៉ោងបង្រៀន"
        }

    start_min = matched_period["start_hour"] * 60 + matched_period["start_min"]
    min_into_period = cur_minutes - start_min
    sec_into_period = min_into_period * 60 + current_dt.second

    sec_remaining_first_30 = max(0, 1800 - sec_into_period)
    sec_remaining_period = max(0, 3600 - sec_into_period)

    if min_into_period < 30:
        phase = "first_30"
        phase_kh = "ចន្លោះពេលធម្មតា (កែប្រែបានច្រើនដង)"
    else:
        phase = "second_30"
        phase_kh = "ចន្លោះពេលបន្ថែម (បញ្ចូលបានតែ ១ ដងគត់)"

    matched_period["phase"] = phase
    matched_period["phase_kh"] = phase_kh
    matched_period["sec_remaining_first_30"] = sec_remaining_first_30
    matched_period["sec_remaining_period"] = sec_remaining_period
    matched_period["min_into_period"] = min_into_period

    return {
        "has_active_period": True,
        "period": matched_period,
        "day": day_info,
        "current_time_str": current_dt.strftime("%H:%M:%S"),
        "current_date_str": current_dt.strftime("%Y-%m-%d"),
        "current_date_display": current_dt.strftime("%d/%m/%Y"),
        "sec_until_next_hour": sec_until_next_hour,
        "phase": phase,
        "phase_kh": phase_kh
    }


def get_current_teaching_slot(teacher_id, current_dt=None):
    """
    ស្វែងរកម៉ោងបង្រៀនរបស់គ្រូនៅម៉ោង និងថ្ងៃបច្ចុប្បន្ន
    """
    period_info = get_current_period_info(current_dt)
    if not period_info["has_active_period"]:
        return {
            "period_info": period_info,
            "slot": None,
            "has_slot": False,
            "reason": "ក្រៅម៉ោងបង្រៀនផ្លូវការ"
        }

    p = period_info["period"]
    day_code = period_info["day"]["code"]
    db_period_num = p["db_period_num"]

    conn = get_db_connection()
    slot = conn.execute("""
        SELECT ts.*, COALESCE(c.id, ts.class_id) as class_db_id, c.class_name, c.room_number as class_room,
               (SELECT COUNT(*) FROM students s WHERE s.class_id = COALESCE(c.id, ts.class_id) AND s.status = 'Active') as student_count
        FROM timetable_slots ts
        LEFT JOIN classes c ON ts.class_id = c.id OR c.class_name LIKE '%' || ts.class_code || '%'
        WHERE ts.teacher_id = ? AND ts.day_code = ? AND ts.period_num = ?
        LIMIT 1
    """, (teacher_id, day_code, db_period_num)).fetchone()
    conn.close()

    if not slot:
        return {
            "period_info": period_info,
            "slot": None,
            "has_slot": False,
            "reason": f"លោកគ្រូ-អ្នកគ្រូ មិនមានម៉ោងបង្រៀនក្នុង {p['label']} នេះទេ"
        }

    slot_dict = dict(slot)
    audit = get_attendance_submission_audit(
        slot_dict["class_db_id"],
        period_info["current_date_str"],
        p["shift"],
        f"Session {db_period_num}"
    )

    return {
        "period_info": period_info,
        "slot": slot_dict,
        "has_slot": True,
        "audit": audit
    }


def get_attendance_submission_audit(class_id, date_str, shift, period_str):
    """ទាញយកព័ត៌មានកំណត់ត្រាដាក់ស្នើវត្តមាន"""
    conn = get_db_connection()
    row = conn.execute("""
        SELECT COUNT(*) as submission_count,
               MIN(submitted_at) as first_submitted_at,
               MAX(submitted_at) as last_submitted_at,
               MAX(phase) as last_phase
        FROM attendance_audit_logs
        WHERE class_id = ? AND date = ? AND shift = ? AND (period = ? OR period = 'Daily')
    """, (class_id, date_str, shift, str(period_str))).fetchone()
    conn.close()

    if row and row["submission_count"] and row["submission_count"] > 0:
        return dict(row)

    # Check if student_attendance already has records for this slot
    conn = get_db_connection()
    marked = conn.execute("""
        SELECT COUNT(*) as cnt, MIN(created_at) as first_time
        FROM student_attendance
        WHERE class_id = ? AND date = ? AND shift = ? AND (period = ? OR period = 'Daily')
    """, (class_id, date_str, shift, str(period_str))).fetchone()
    conn.close()

    if marked and marked["cnt"] and marked["cnt"] > 0:
        return {
            "submission_count": 1,
            "first_submitted_at": marked["first_time"],
            "last_submitted_at": marked["first_time"],
            "last_phase": "first_30"
        }

    return {
        "submission_count": 0,
        "first_submitted_at": None,
        "last_submitted_at": None,
        "last_phase": None
    }


def record_attendance_audit(class_id, date_str, shift, period_str, period_num, teacher_id, phase, ip_address=None, notes=None):
    """កត់ត្រា Log នៃការដាក់ស្នើវត្តមាន"""
    conn = get_db_connection()
    cur_cnt = conn.execute("""
        SELECT COUNT(*) as c FROM attendance_audit_logs
        WHERE class_id = ? AND date = ? AND shift = ? AND period = ?
    """, (class_id, date_str, shift, str(period_str))).fetchone()["c"]

    new_count = cur_cnt + 1
    conn.execute("""
        INSERT INTO attendance_audit_logs (date, class_id, shift, period, period_num, teacher_id, phase, submission_count, ip_address, notes)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (date_str, class_id, shift, str(period_str), period_num, teacher_id, phase, new_count, ip_address, notes))
    conn.commit()
    conn.close()
    return new_count


def check_submission_window_rule(class_id, date_str, shift, period_str, teacher_id=None, is_admin=False, current_dt=None):
    """
    ផ្ទៀងផ្ទាត់លក្ខខណ្ឌម៉ោងស្រង់វត្តមាន៖
    - បើជា Admin: អនុញ្ញាតដោយសេរី (Administrative override)
    - បើជា Teacher:
      1. ត្រូវតែជាថ្ងៃ និងម៉ោងបច្ចុប្បន្ន (បដិសេធកាលបរិច្ឆេទខុស ឬម៉ោងមុន)
      2. គ្រូត្រូវតែមានកាលវិភាគបង្រៀនថ្នាក់នេះពិតប្រាកដ
      3. ក្នុង ៣០ នាទីដំបូង (ឧ. ៧:០០ - ៧:៣០)៖ អនុញ្ញាតច្រើនដង
      4. ក្នុង ៣០ នាទីក្រោយ (ឧ. ៧:៣០ - ៨:០០)៖ អនុញ្ញាតតែ ១ ដងគត់ ប្រសិនបើមិនទាន់បានដាក់ស្នើ
      5. ផុតម៉ោង (ឧ. >= ៨:០០)៖ បដិសេធដាច់ខាត
    """
    if is_admin:
        return True, 200, "អនុញ្ញាត (Administrator)", "admin", 0

    if current_dt is None:
        current_dt = datetime.now()

    period_info = get_current_period_info(current_dt)

    # 1. Check Date
    today_date = current_dt.strftime("%Y-%m-%d")
    if date_str != today_date:
        return False, 403, "លោកគ្រូ-អ្នកគ្រូ មិនអាចស្រង់វត្តមានខុសពីកាលបរិច្ឆេទថ្ងៃនេះបានឡើយ!", "invalid_date", 0

    # 2. Check if currently in an active teaching period
    if not period_info["has_active_period"]:
        return False, 403, "បច្ចុប្បន្នស្ថិតនៅក្រៅម៉ោងបង្រៀនផ្លូវការ! ប្រព័ន្ធមិនអនុញ្ញាតឱ្យបញ្ចូលទិន្នន័យឡើយ។", "outside_hours", 0

    active_p = period_info["period"]

    # 3. Check Shift
    if shift != active_p["shift"]:
        return False, 403, f"វេនសិក្សាមិនត្រឹមត្រូវទេ! ពេលនេះជាវេន {active_p['shift_kh']}។", "wrong_shift", 0

    # 4. Check Period match
    req_period_num = None
    if "Session " in str(period_str):
        try:
            req_period_num = int(str(period_str).replace("Session ", "").strip())
        except Exception:
            pass
    elif "ម៉ោងទី " in str(period_str):
        try:
            num_part = int(''.join(filter(str.isdigit, str(period_str))))
            req_period_num = num_part if shift == "Morning" else num_part + 4
        except Exception:
            pass
    elif str(period_str).isdigit():
        req_period_num = int(period_str)

    if req_period_num and req_period_num != active_p["db_period_num"]:
        return False, 403, f"ផុតកំណត់ម៉ោងស្រង់វត្តមានហើយ! ប្រព័ន្ធបដិសេធការបញ្ជូនទិន្នន័យម៉ោងមុន។ សូមស្រង់វត្តមានក្នុងម៉ោងបច្ចុប្បន្ន ({active_p['label']}) ជំនួសវិញ។", "expired_period", 0

    # 5. Check Teacher Timetable slot
    if teacher_id:
        day_code = period_info["day"]["code"]
        conn = get_db_connection()
        slot = conn.execute("""
            SELECT id FROM timetable_slots
            WHERE teacher_id = ? AND day_code = ? AND period_num = ?
              AND (class_id = ? OR class_code = (SELECT REPLACE(class_name, 'ថ្នាក់ទី ', '') FROM classes WHERE id = ?))
            LIMIT 1
        """, (teacher_id, day_code, active_p["db_period_num"], class_id, class_id)).fetchone()
        conn.close()

        if not slot:
            return False, 403, f"លោកគ្រូ-អ្នកគ្រូ មិនមានកាលវិភាគបង្រៀននៅថ្នាក់នេះក្នុងម៉ោង {active_p['short_label']} ទេ!", "unauthorized_slot", 0

    # 6. Check Time Phase & Submission Count
    period_key = f"Session {active_p['db_period_num']}"
    audit = get_attendance_submission_audit(class_id, today_date, shift, period_key)
    sub_count = audit.get("submission_count", 0)

    if active_p["phase"] == "first_30":
        # First 30 mins: allowed multiple times
        return True, 200, "អនុញ្ញាតក្នុងចន្លោះ ៣០ នាទីដំបូង", "first_30", sub_count
    else:
        # Second 30 mins (30 - 60 mins): allowed ONCE if not submitted yet
        if sub_count == 0:
            return True, 200, "អនុញ្ញាតបញ្ចូលបាន ១ ដងគត់ក្នុងចន្លោះពេលបន្ថែម (៣០ នាទីចុងក្រោយ)", "second_30", sub_count
        else:
            return False, 403, f"ផុតម៉ោងអនុញ្ញាតកែប្រែ (៣០ នាទីដំបូង) ហើយ! លោកគ្រូ-អ្នកគ្រូ បានដាក់ស្នើរួចហើយ ({sub_count} ដង) មិនអាចកែប្រែបានទៀតទេ។", "locked_second_30", sub_count


if __name__ == "__main__":
    init_db()
    init_default_users()
