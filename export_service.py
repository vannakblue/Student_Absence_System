"""
ប្រព័ន្ធគ្រប់គ្រងអវត្តមានសិស្ស-គ្រូ (Student & Teacher Absence Management System)
Excel Export Service (openpyxl)
"""

import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from database import (
    get_student_report_data, get_teacher_report_data, 
    get_class_by_id, get_setting
)


def get_khmer_font(name="Kantumruy Pro", size=11, bold=False, italic=False, color="000000"):
    return Font(name=name, size=size, bold=bold, italic=italic, color=color)


def get_thin_border():
    thin = Side(border_style="thin", color="CBD5E1")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def export_student_attendance_excel(class_id=None, start_date=None, end_date=None, output_path=None):
    """
    Export របាយការណ៍អវត្តមានសិស្សជា Excel
    """
    if not output_path:
        filename = f"របាយការណ៍វត្តមានសិស្ស_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scratch", filename)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

    data = get_student_report_data(class_id=class_id, start_date=start_date, end_date=end_date)
    school_name = get_setting("school_name_kh", "វិទ្យាល័យ ហ៊ុន សែន")
    class_info = get_class_by_id(class_id) if class_id else None
    class_label = class_info["class_name"] if class_info else "ថ្នាក់ទាំងអស់"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "របាយការណ៍វត្តមានសិស្ស"
    ws.views.sheetView[0].showGridLines = True

    # Title Headers
    ws.merge_cells("A1:H1")
    ws["A1"] = "ព្រះរាជាណាចក្រកម្ពុជា"
    ws["A1"].font = get_khmer_font(size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = "ជាតិ សាសនា ព្រះមហាក្សត្រ"
    ws["A2"].font = get_khmer_font(size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A4"] = school_name
    ws["A4"].font = get_khmer_font(size=11, bold=True)

    ws["A5"] = f"របាយការណ៍ស្រង់វត្តមានសិស្ស៖ {class_label}"
    ws["A5"].font = get_khmer_font(size=13, bold=True, color="1E3A8A")

    date_str = f"កាលបរិច្ឆេទ៖ {start_date or 'ដើមឆ្នាំ'} ដល់ {end_date or 'បច្ចុប្បន្ន'}"
    ws["A6"] = date_str
    ws["A6"].font = get_khmer_font(size=10, italic=True)

    # Table Header Row
    headers = [
        ("ល.រ", 6),
        ("អត្តលេខ", 14),
        ("គោត្តនាម-នាម", 24),
        ("ភេទ", 8),
        ("ថ្នាក់", 14),
        ("មានច្បាប់ (P)", 14),
        ("ឥតច្បាប់ (A)", 14),
        ("មកយឺត (L)", 14),
        ("វត្តមាន (Pr)", 14),
        ("សរុបអវត្តមាន", 14),
    ]

    header_row = 8
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(name="Kantumruy Pro", size=10, bold=True, color="FFFFFF")

    for col_idx, (h_title, col_width) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = get_thin_border()
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    # Data Rows
    thin_border = get_thin_border()
    current_row = header_row + 1

    for idx, s in enumerate(data, start=1):
        absent_total = (s["permission_count"] or 0) + (s["absent_count"] or 0)
        row_vals = [
            idx,
            s["student_code"],
            s["full_name_kh"],
            s["gender"],
            s["class_name"],
            s["permission_count"] or 0,
            s["absent_count"] or 0,
            s["late_count"] or 0,
            s["present_count"] or 0,
            absent_total
        ]

        row_fill = PatternFill(start_color="F8FAFC" if idx % 2 == 0 else "FFFFFF", fill_type="solid")

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=current_row, column=c_idx, value=val)
            cell.font = get_khmer_font(size=10)
            cell.border = thin_border
            cell.fill = row_fill
            if c_idx in [1, 2, 4, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 3:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        current_row += 1

    # Official Footer Approval & Signature
    current_row += 2
    sign_col = 8
    ws.cell(row=current_row, column=sign_col, value="ថ្ងៃ............ ខែ............ ឆ្នាំ............ ព.ស. ២៥៧០").font = get_khmer_font(size=9, italic=True)
    current_row += 1
    ws.cell(row=current_row, column=sign_col, value=f"ធ្វើនៅកំពង់កន្ទួត, ថ្ងៃទី {datetime.now().strftime('%d-%m-%Y')}").font = get_khmer_font(size=9, italic=True)
    current_row += 1
    ws.cell(row=current_row, column=sign_col, value="បានឃើញ និងឯកភាព").font = get_khmer_font(size=10, bold=True)
    current_row += 1
    ws.cell(row=current_row, column=sign_col, value="នាយកវិទ្យាល័យ").font = get_khmer_font(size=11, bold=True)
    current_row += 3
    principal_name = get_setting("principal_name", "លោកបណ្ឌិត សុខ ចាន់ថន")
    ws.cell(row=current_row, column=sign_col, value=principal_name).font = get_khmer_font(size=11, bold=True, color="1E3A8A")

    wb.save(output_path)
    return output_path


def export_teacher_attendance_excel(start_date=None, end_date=None, output_path=None):
    """
    Export របាយការណ៍វត្តមានគ្រូបង្រៀនជា Excel
    """
    if not output_path:
        filename = f"របាយការណ៍វត្តមានគ្រូ_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scratch", filename)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

    data = get_teacher_report_data(start_date=start_date, end_date=end_date)
    school_name = get_setting("school_name_kh", "វិទ្យាល័យ ហ៊ុន សែន")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "របាយការណ៍វត្តមានគ្រូ"
    ws.views.sheetView[0].showGridLines = True

    # Title Headers
    ws.merge_cells("A1:I1")
    ws["A1"] = "ព្រះរាជាណាចក្រកម្ពុជា"
    ws["A1"].font = get_khmer_font(size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:I2")
    ws["A2"] = "ជាតិ សាសនា ព្រះមហាក្សត្រ"
    ws["A2"].font = get_khmer_font(size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A4"] = school_name
    ws["A4"].font = get_khmer_font(size=11, bold=True)

    ws["A5"] = "របាយការណ៍បូកសរុបវត្តមានគ្រូបង្រៀន"
    ws["A5"].font = get_khmer_font(size=13, bold=True, color="047857")

    date_str = f"កាលបរិច្ឆេទ៖ {start_date or 'ដើមខែ'} ដល់ {end_date or 'បច្ចុប្បន្ន'}"
    ws["A6"] = date_str
    ws["A6"].font = get_khmer_font(size=10, italic=True)

    headers = [
        ("ល.រ", 6),
        ("កូដគ្រូ", 12),
        ("គោត្តនាម-នាម", 24),
        ("ភេទ", 8),
        ("មុខវិជ្ជា", 16),
        ("មានច្បាប់ (P)", 14),
        ("ឥតច្បាប់ (A)", 14),
        ("មកយឺត (L)", 14),
        ("វត្តមាន (Pr)", 14),
    ]

    header_row = 8
    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_font = Font(name="Kantumruy Pro", size=10, bold=True, color="FFFFFF")

    for col_idx, (h_title, col_width) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = get_thin_border()
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    thin_border = get_thin_border()
    current_row = header_row + 1

    for idx, t in enumerate(data, start=1):
        row_vals = [
            idx,
            t["teacher_code"],
            t["full_name_kh"],
            t["gender"],
            t["subject"],
            t["permission_count"] or 0,
            t["absent_count"] or 0,
            t["late_count"] or 0,
            t["present_count"] or 0,
        ]

        row_fill = PatternFill(start_color="F8FAFC" if idx % 2 == 0 else "FFFFFF", fill_type="solid")

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=current_row, column=c_idx, value=val)
            cell.font = get_khmer_font(size=10)
            cell.border = thin_border
            cell.fill = row_fill
            if c_idx in [1, 2, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [3, 5]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        current_row += 1

    # Official Footer Approval & Signature
    current_row += 2
    sign_col = 7
    ws.cell(row=current_row, column=sign_col, value="ថ្ងៃ............ ខែ............ ឆ្នាំ............ ព.ស. ២៥៧០").font = get_khmer_font(size=9, italic=True)
    current_row += 1
    ws.cell(row=current_row, column=sign_col, value=f"ធ្វើនៅកំពង់កន្ទួត, ថ្ងៃទី {datetime.now().strftime('%d-%m-%Y')}").font = get_khmer_font(size=9, italic=True)
    current_row += 1
    ws.cell(row=current_row, column=sign_col, value="បានឃើញ និងឯកភាព").font = get_khmer_font(size=10, bold=True)
    current_row += 1
    ws.cell(row=current_row, column=sign_col, value="នាយកវិទ្យាល័យ").font = get_khmer_font(size=11, bold=True)
    current_row += 3
    principal_name = get_setting("principal_name", "លោកបណ្ឌិត សុខ ចាន់ថន")
    ws.cell(row=current_row, column=sign_col, value=principal_name).font = get_khmer_font(size=11, bold=True, color="065F46")

    wb.save(output_path)
    return output_path


def export_users_excel(output_path=None):
    """
    Export បញ្ជីគណនី Username និង Password សម្រាប់ Admin ចែកជូនគ្រូបង្រៀន
    """
    from database import get_all_users

    if not output_path:
        filename = f"បញ្ជីគណនីគ្រូបង្រៀន_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scratch", filename)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

    users = get_all_users()
    school_name = get_setting("school_name_kh", "វិទ្យាល័យ ហ៊ុន សែន កំពង់កន្ទួត")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "គណនីគ្រូបង្រៀន"
    ws.views.sheetView[0].showGridLines = True

    # Title Headers
    ws.merge_cells("A1:G1")
    ws["A1"] = school_name
    ws["A1"].font = get_khmer_font(size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = "បញ្ជីឈ្មោះ និងគណនីចូលប្រើប្រាស់ប្រព័ន្ធ (User Accounts)"
    ws["A2"].font = get_khmer_font(size=12, bold=True, color="1E40AF")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A4"] = f"កាលបរិច្ឆេទបង្កើត៖ {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A4"].font = get_khmer_font(size=10, italic=True)

    headers = [
        ("ល.រ", 6),
        ("Username (កូដគ្រូ)", 18),
        ("គោត្តនាម និងនាម", 24),
        ("មុខវិជ្ជាឯកទេស", 18),
        ("លេខទូរស័ព្ទ", 16),
        ("Password លំនាំដើម", 18),
        ("តួនាទី (Role)", 14),
    ]

    header_row = 6
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Kantumruy Pro", size=10, bold=True, color="FFFFFF")

    for col_idx, (h_title, col_width) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = get_thin_border()
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    thin_border = get_thin_border()
    current_row = header_row + 1

    for idx, u in enumerate(users, start=1):
        row_vals = [
            idx,
            u["username"],
            u["full_name_kh"],
            u.get("subject") or "—",
            u.get("phone") or "—",
            u.get("plain_password_hint") or "******",
            "Admin" if u["role"] == "admin" else "គ្រូបង្រៀន (Teacher)"
        ]

        row_fill = PatternFill(start_color="F8FAFC" if idx % 2 == 0 else "FFFFFF", fill_type="solid")

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=current_row, column=c_idx, value=val)
            cell.font = get_khmer_font(size=10, bold=(c_idx in [2, 6]))
            cell.border = thin_border
            cell.fill = row_fill
            if c_idx in [1, 2, 5, 6, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    wb.save(output_path)
    return output_path


def export_timetable_template_excel(output_path=None):
    """
    បង្កើតឯកសារ Excel Template សម្រាប់ Admin បញ្ចូលកាលវិភាគរួម (Master School Timetable Template)
    មាន Sheet ទី១ ជាតារាងកាលវិភាគ និង Sheet ទី២ ជាតារាងយោង (References)
    """
    import database as db

    if not output_path:
        filename = f"ទម្រង់គំរូ_កាលវិភាគរួម_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scratch", filename)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

    school_name = db.get_setting("school_name_kh", "វិទ្យាល័យ ហ៊ុន សែន")
    classes = db.get_classes()
    teachers = db.get_teachers(active_only=True)
    subjects = db.get_all_subjects()

    wb = openpyxl.Workbook()
    
    # Sheet 1: Master Timetable Input Sheet
    ws1 = wb.active
    ws1.title = "កាលវិភាគរួម"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:G1")
    ws1["A1"] = school_name
    ws1["A1"].font = get_khmer_font(size=14, bold=True)
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws1.merge_cells("A2:G2")
    ws1["A2"] = "ទម្រង់បញ្ចូលកាលវិភាគរួមក្នុងសាលា (Master School Timetable Import Template)"
    ws1["A2"].font = get_khmer_font(size=12, bold=True, color="1E40AF")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws1.merge_cells("A3:G3")
    ws1["A3"] = "ការណែនាំ៖ បំពេញកាលវិភាគតាមជួរដេកនីមួយៗ (Class, Day, Period, Subject, Teacher Code/Name, Room)។ មើលសន្លឹកកិច្ចការ 'តារាងយោង' សម្រាប់កូដគ្រូ និងថ្នាក់។"
    ws1["A3"].font = get_khmer_font(size=9, italic=True, color="475569")
    ws1["A3"].alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        ("ល.រ", 6),
        ("ថ្នាក់រៀន (Class Code)*", 18),
        ("ថ្ងៃបង្រៀន (Day Code)*", 18),
        ("ម៉ោងទី (Period 1-8)*", 18),
        ("មុខវិជ្ជា (Subject Name)*", 24),
        ("គ្រូបង្រៀន (Teacher Code / Name)*", 28),
        ("បន្ទប់សិក្សា (Room)", 16),
    ]

    header_row = 5
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Kantumruy Pro", size=10, bold=True, color="FFFFFF")
    thin_border = get_thin_border()

    for col_idx, (h_title, col_width) in enumerate(headers, start=1):
        cell = ws1.cell(row=header_row, column=col_idx, value=h_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = col_width

    # Sample demo rows for reference
    sample_rows = [
        (1, "7A", "ច", 1, "ភាសាខ្មែរ", teachers[0]["teacher_code"] if teachers else "T001", "101"),
        (2, "7A", "ច", 2, "ភាសាខ្មែរ", teachers[0]["teacher_code"] if teachers else "T001", "101"),
        (3, "7A", "ច", 3, "គណិតវិទ្យា", teachers[1]["teacher_code"] if len(teachers) > 1 else "T002", "101"),
        (4, "7A", "ច", 4, "រូបវិទ្យា", teachers[2]["teacher_code"] if len(teachers) > 2 else "T003", "101"),
        (5, "10A", "ច", 5, "គីមីវិទ្យា", teachers[3]["teacher_code"] if len(teachers) > 3 else "T004", "201"),
    ]

    for r_idx, row_vals in enumerate(sample_rows, start=header_row + 1):
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws1.cell(row=r_idx, column=c_idx, value=val)
            cell.font = get_khmer_font(size=10)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Sheet 2: Reference Lookup Sheet
    ws2 = wb.create_sheet(title="តារាងយោង")
    ws2.views.sheetView[0].showGridLines = True

    ws2["A1"] = "តារាងយោងកាលវិភាគ (Timetable References)"
    ws2["A1"].font = get_khmer_font(size=13, bold=True, color="1E40AF")

    # Days reference
    ws2["A3"] = "១. កូដថ្ងៃ (Day Codes)"
    ws2["A3"].font = get_khmer_font(size=11, bold=True)
    ws2["A4"] = "កូដថ្ងៃ"
    ws2["B4"] = "ឈ្មោះថ្ងៃ"
    ws2["A4"].font = ws2["B4"].font = get_khmer_font(size=10, bold=True)
    days_data = [
        ("ច", "ចន្ទ (Monday)"),
        ("អ", "អង្គារ (Tuesday)"),
        ("ព", "ពុធ (Wednesday)"),
        ("ព្រ", "ព្រហស្បតិ៍ (Thursday)"),
        ("សុ", "សុក្រ (Friday)"),
        ("ស", "សៅរ៍ (Saturday)"),
    ]
    for idx, (d_code, d_name) in enumerate(days_data, start=5):
        ws2.cell(row=idx, column=1, value=d_code).font = get_khmer_font(size=10, bold=True)
        ws2.cell(row=idx, column=2, value=d_name).font = get_khmer_font(size=10)

    # Periods reference
    ws2["D3"] = "២. ម៉ោងសិក្សា (Periods)"
    ws2["D3"].font = get_khmer_font(size=11, bold=True)
    ws2["D4"] = "លេខម៉ោង"
    ws2["E4"] = "វេនសិក្សា"
    ws2["F4"] = "ចន្លោះម៉ោង"
    ws2["D4"].font = ws2["E4"].font = ws2["F4"].font = get_khmer_font(size=10, bold=True)
    periods_data = [
        (1, "ពេលព្រឹក (Morning)", "07:00 - 08:00 (ម៉ោងទី ១)"),
        (2, "ពេលព្រឹក (Morning)", "08:00 - 09:00 (ម៉ោងទី ២)"),
        (3, "ពេលព្រឹក (Morning)", "09:00 - 10:00 (ម៉ោងទី ៣)"),
        (4, "ពេលព្រឹក (Morning)", "10:00 - 11:00 (ម៉ោងទី ៤)"),
        (5, "ពេលរសៀល (Afternoon)", "13:00 - 14:00 (ម៉ោងទី ១ រសៀល)"),
        (6, "ពេលរសៀល (Afternoon)", "14:00 - 15:00 (ម៉ោងទី ២ រសៀល)"),
        (7, "ពេលរសៀល (Afternoon)", "15:00 - 16:00 (ម៉ោងទី ៣ រសៀល)"),
        (8, "ពេលរសៀល (Afternoon)", "16:00 - 17:00 (ម៉ោងទី ៤ រសៀល)"),
    ]
    for idx, (p_num, p_shift, p_time) in enumerate(periods_data, start=5):
        ws2.cell(row=idx, column=4, value=p_num).font = get_khmer_font(size=10, bold=True)
        ws2.cell(row=idx, column=5, value=p_shift).font = get_khmer_font(size=10)
        ws2.cell(row=idx, column=6, value=p_time).font = get_khmer_font(size=10)

    # Classes reference
    ws2["H3"] = "៣. បញ្ជីថ្នាក់រៀន (Classes)"
    ws2["H3"].font = get_khmer_font(size=11, bold=True)
    ws2["H4"] = "កូដថ្នាក់"
    ws2["I4"] = "ឈ្មោះថ្នាក់"
    ws2["H4"].font = ws2["I4"].font = get_khmer_font(size=10, bold=True)
    for idx, c in enumerate(classes, start=5):
        c_code = c["class_name"].replace("ថ្នាក់ទី ", "").strip()
        ws2.cell(row=idx, column=8, value=c_code).font = get_khmer_font(size=10, bold=True)
        ws2.cell(row=idx, column=9, value=c["class_name"]).font = get_khmer_font(size=10)

    # Teachers reference
    ws2["K3"] = "៤. បញ្ជីគ្រូបង្រៀន (Teachers)"
    ws2["K3"].font = get_khmer_font(size=11, bold=True)
    ws2["K4"] = "កូដគ្រូ (Code)"
    ws2["L4"] = "ឈ្មោះគ្រូបង្រៀន"
    ws2["M4"] = "មុខវិជ្ជាឯកទេស"
    ws2["K4"].font = ws2["L4"].font = ws2["M4"].font = get_khmer_font(size=10, bold=True)
    for idx, t in enumerate(teachers, start=5):
        ws2.cell(row=idx, column=11, value=t["teacher_code"]).font = get_khmer_font(size=10, bold=True)
        ws2.cell(row=idx, column=12, value=t["full_name_kh"]).font = get_khmer_font(size=10)
        ws2.cell(row=idx, column=13, value=t["subject"]).font = get_khmer_font(size=10)

    ws2.column_dimensions["A"].width = 10
    ws2.column_dimensions["B"].width = 22
    ws2.column_dimensions["D"].width = 10
    ws2.column_dimensions["E"].width = 22
    ws2.column_dimensions["F"].width = 28
    ws2.column_dimensions["H"].width = 12
    ws2.column_dimensions["I"].width = 20
    ws2.column_dimensions["K"].width = 18
    ws2.column_dimensions["L"].width = 24
    ws2.column_dimensions["M"].width = 18

    wb.save(output_path)
    return output_path


def export_master_timetable_excel(output_path=None):
    """
    Export កាលវិភាគរួមរបស់សាលាទាំងមូលចេញជា Excel
    ដើម្បីឱ្យ Admin អាចទាញយកទៅមើល កែប្រែ ឬ Upload ត្រឡប់មកវិញ
    """
    import database as db

    if not output_path:
        filename = f"កាលវិភាគរួម_សាលា_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scratch", filename)
        os.makedirs(os.path.dirname(output_path), exist_ok=True)

    school_name = db.get_setting("school_name_kh", "វិទ្យាល័យ ហ៊ុន សែន")

    conn = db.get_db_connection()
    slots = conn.execute("""
        SELECT ts.*, t.full_name_kh as teacher_full_name, t.teacher_code as real_teacher_code
        FROM timetable_slots ts
        LEFT JOIN teachers t ON ts.teacher_id = t.id
        ORDER BY 
            ts.class_code ASC,
            CASE ts.day_code
                WHEN 'ច' THEN 1
                WHEN 'អ' THEN 2
                WHEN 'ព' THEN 3
                WHEN 'ព្រ' THEN 4
                WHEN 'សុ' THEN 5
                WHEN 'ស' THEN 6
                ELSE 7
            END,
            ts.period_num ASC
    """).fetchall()
    conn.close()

    # Generate using template layout
    wb = openpyxl.load_workbook(export_timetable_template_excel())
    ws1 = wb["កាលវិភាគរួម"]

    # Clear sample rows (from row 6 downward)
    max_r = ws1.max_row
    if max_r >= 6:
        for r in range(6, max_r + 1):
            for c in range(1, 8):
                ws1.cell(row=r, column=c, value=None)

    thin_border = get_thin_border()
    current_row = 6

    for idx, s in enumerate(slots, start=1):
        s_dict = dict(s)
        teacher_display = s_dict.get("real_teacher_code") or s_dict.get("teacher_code") or s_dict.get("teacher_full_name") or (s_dict.get("teacher_name") or "")
        row_vals = [
            idx,
            s_dict.get("class_code", ""),
            s_dict.get("day_code", ""),
            s_dict.get("period_num", ""),
            s_dict.get("subject_name", ""),
            teacher_display,
            s_dict.get("room_number") or ""
        ]
        row_fill = PatternFill(start_color="F8FAFC" if idx % 2 == 0 else "FFFFFF", fill_type="solid")

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws1.cell(row=current_row, column=c_idx, value=val)
            cell.font = get_khmer_font(size=10)
            cell.border = thin_border
            cell.fill = row_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        current_row += 1

    wb.save(output_path)
    return output_path

